import io
import calendar
from datetime import date, timedelta

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from Controllers.thong_ke_controller import ThongKeController
from Views.layout import (
    hien_thi_header,
    hien_thi_the,
    hien_thi_tieu_de_section,
)


def _tao_cac_tuan_trong_thang(nam, thang):
    ngay_dau_thang = date(nam, thang, 1)

    ngay_cuoi_thang = date(
        nam,
        thang,
        calendar.monthrange(nam, thang)[1],
    )

    thu_hai_dau = (
        ngay_dau_thang
        - timedelta(days=ngay_dau_thang.weekday())
    )

    thu_bay_cuoi = (
        ngay_cuoi_thang
        + timedelta(days=5 - ngay_cuoi_thang.weekday())
    )

    danh_sach = []
    ngay_bat_dau = thu_hai_dau
    so_tuan = 1

    while ngay_bat_dau <= thu_bay_cuoi:
        ngay_ket_thuc = ngay_bat_dau + timedelta(days=5)

        if (
            ngay_ket_thuc >= ngay_dau_thang
            and ngay_bat_dau <= ngay_cuoi_thang
        ):
            danh_sach.append(
                (
                    so_tuan,
                    ngay_bat_dau,
                    ngay_ket_thuc,
                )
            )

            so_tuan += 1

        ngay_bat_dau += timedelta(days=7)

    return danh_sach


def _tao_dataframe(rows):
    du_lieu = []

    for stt, row in enumerate(rows, start=1):
        du_lieu.append({
            "STT": stt,
            "Họ tên": str(row[1] or ""),
            "Bộ phận": str(
                row[2] or "Chưa phân bộ phận"
            ),
            "Số ngày ăn": int(row[3] or 0),
            "Số suất": int(row[4] or 0),
            "Tiền cơm": float(row[5] or 0),
        })

    return pd.DataFrame(
        du_lieu,
        columns=[
            "STT",
            "Họ tên",
            "Bộ phận",
            "Số ngày ăn",
            "Số suất",
            "Tiền cơm",
        ],
    )


def _xuat_excel(
    df,
    tu_ngay,
    den_ngay,
    kieu,
    bo_phan,
    summary,
):
    output = io.BytesIO()

    wb = Workbook()
    ws = wb.active
    ws.title = "Thống kê suất ăn"

    mau_tieu_de = "E8DDEB"
    mau_header = "D9EAF7"
    mau_tong = "F3E5F5"

    vien = Side(
        style="thin",
        color="D9D9D9",
    )

    ws.merge_cells("A1:F1")
    ws["A1"] = "BÁO CÁO THỐNG KÊ SUẤT ĂN"
    ws["A1"].font = Font(
        bold=True,
        size=16,
        name="Times New Roman",
    )
    ws["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )
    ws["A1"].fill = PatternFill(
        "solid",
        fgColor=mau_tieu_de,
    )
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:F2")
    ws["A2"] = (
        f"Kỳ thống kê: {kieu} | "
        f"Từ {tu_ngay.strftime('%d/%m/%Y')} "
        f"đến {den_ngay.strftime('%d/%m/%Y')}"
    )
    ws["A2"].font = Font(
        name="Times New Roman",
        size=13,
    )
    ws["A2"].alignment = Alignment(
        horizontal="center",
    )

    ws.merge_cells("A3:F3")
    ws["A3"] = f"Bộ phận: {bo_phan}"
    ws["A3"].font = Font(
        name="Times New Roman",
        size=13,
    )
    ws["A3"].alignment = Alignment(
        horizontal="center",
    )

    thong_tin = [
        (
            "Tổng người ăn",
            int(summary["tong_nguoi"]),
        ),
        (
            "Tổng ngày ăn",
            int(summary["tong_ngay_an"]),
        ),
        (
            "Tổng suất",
            int(summary["tong_suat"]),
        ),
        (
            "Tổng tiền",
            float(summary["tong_tien"]),
        ),
    ]

    for cot, (ten, gia_tri) in enumerate(
        thong_tin,
        start=1,
    ):
        ws.cell(5, cot, ten)
        ws.cell(6, cot, gia_tri)

        ws.cell(5, cot).font = Font(
            bold=True,
            name="Times New Roman",
            size=13,
        )

        ws.cell(6, cot).font = Font(
            name="Times New Roman",
            size=13,
        )

        ws.cell(5, cot).fill = PatternFill(
            "solid",
            fgColor=mau_header,
        )

        ws.cell(5, cot).alignment = Alignment(
            horizontal="center",
        )

        ws.cell(6, cot).alignment = Alignment(
            horizontal="center",
        )

        ws.cell(5, cot).border = Border(
            bottom=vien,
        )

        ws.cell(6, cot).border = Border(
            bottom=vien,
        )

    ws["D6"].number_format = '#,##0 "đ"'

    dong_header = 8

    for cot, ten_cot in enumerate(
        df.columns,
        start=1,
    ):
        cell = ws.cell(
            dong_header,
            cot,
            ten_cot,
        )

        cell.font = Font(
            bold=True,
            name="Times New Roman",
            size=13,
        )

        cell.fill = PatternFill(
            "solid",
            fgColor=mau_header,
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        cell.border = Border(
            top=vien,
            bottom=vien,
            left=vien,
            right=vien,
        )

    for dong, record in enumerate(
        df.itertuples(index=False),
        start=dong_header + 1,
    ):
        for cot, gia_tri in enumerate(
            record,
            start=1,
        ):
            cell = ws.cell(
                dong,
                cot,
                gia_tri,
            )

            cell.font = Font(
                name="Times New Roman",
                size=13,
            )

            cell.border = Border(
                top=vien,
                bottom=vien,
                left=vien,
                right=vien,
            )

            cell.alignment = Alignment(
                vertical="center",
            )

            if cot == 1:
                cell.alignment = Alignment(
                    horizontal="center",
                )

            elif cot in (4, 5):
                cell.alignment = Alignment(
                    horizontal="center",
                )

            elif cot == 6:
                cell.number_format = '#,##0 "đ"'
                cell.alignment = Alignment(
                    horizontal="right",
                )

    dong_tong = (
        dong_header
        + len(df)
        + 1
    )

    ws.cell(
        dong_tong,
        1,
        "TỔNG",
    )

    ws.merge_cells(
        start_row=dong_tong,
        start_column=1,
        end_row=dong_tong,
        end_column=3,
    )

    ws.cell(
        dong_tong,
        4,
        int(summary["tong_ngay_an"]),
    )

    ws.cell(
        dong_tong,
        5,
        int(summary["tong_suat"]),
    )

    ws.cell(
        dong_tong,
        6,
        float(summary["tong_tien"]),
    )

    for cot in range(1, 7):
        cell = ws.cell(
            dong_tong,
            cot,
        )

        cell.font = Font(
            bold=True,
            name="Times New Roman",
            size=13,
        )

        cell.fill = PatternFill(
            "solid",
            fgColor=mau_tong,
        )

        cell.border = Border(
            top=vien,
            bottom=vien,
            left=vien,
            right=vien,
        )

        cell.alignment = Alignment(
            horizontal=(
                "right"
                if cot == 6
                else "center"
            )
        )

    ws.cell(
        dong_tong,
        6,
    ).number_format = '#,##0 "đ"'

    ws.freeze_panes = "A9"

    if not df.empty:
        ws.auto_filter.ref = (
            f"A8:F{dong_header + len(df)}"
        )
    else:
        ws.auto_filter.ref = "A8:F8"

    do_rong = [
        8,
        28,
        24,
        14,
        12,
        18,
    ]

    for i, rong in enumerate(
        do_rong,
        start=1,
    ):
        ws.column_dimensions[
            get_column_letter(i)
        ].width = rong

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1

    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5

    wb.save(output)
    output.seek(0)

    return output.getvalue()


def hien_thi_thong_ke():
    language = st.session_state.get(
        "language",
        "vi",
    )

    text = {
        "vi": {
            "header": "📊 Thống kê",
            "desc": (
                "Theo dõi tình hình suất ăn "
                "theo tuần hoặc theo tháng."
            ),
            "type": "Kiểu thống kê",
            "week": "Theo tuần",
            "month": "Theo tháng",
            "year": "Năm",
            "month_label": "Tháng",
            "week_label": "Tuần",
            "department": "Bộ phận",
            "all": "Tất cả",
            "period": "Kỳ thống kê",
            "overview": "Tổng quan",
            "people": "Người ăn",
            "days": "Ngày ăn",
            "meals": "Tổng suất",
            "money": "Tiền suất ăn",
            "detail": "Bảng thống kê tổng hợp",
            "empty": (
                "🌷 Chưa có dữ liệu trong kỳ này."
            ),
            "export": "📥 Xuất Excel",
        },

        "en": {
            "header": "📊 Statistics",
            "desc": (
                "Track meal activity "
                "by week or month."
            ),
            "type": "Statistics type",
            "week": "By week",
            "month": "By month",
            "year": "Year",
            "month_label": "Month",
            "week_label": "Week",
            "department": "Department",
            "all": "All",
            "period": "Statistics period",
            "overview": "Overview",
            "people": "People eating",
            "days": "Meal days",
            "meals": "Total meals",
            "money": "Meal amount",
            "detail": "Summary table",
            "empty": (
                "🌷 No data for this period."
            ),
            "export": "📥 Export Excel",
        },

        "zh": {
            "header": "📊 统计",
            "desc": "按周或按月查看用餐情况。",
            "type": "统计类型",
            "week": "按周",
            "month": "按月",
            "year": "年份",
            "month_label": "月份",
            "week_label": "星期",
            "department": "部门",
            "all": "全部",
            "period": "统计期间",
            "overview": "概览",
            "people": "用餐人数",
            "days": "用餐天数",
            "meals": "总餐数",
            "money": "餐费",
            "detail": "汇总表",
            "empty": "🌷 此期间暂无数据。",
            "export": "📥 导出 Excel",
        },
    }.get(
        language,
        {},
    )

    hien_thi_header(
        text["header"],
        text["desc"],
    )

    hien_thi_tieu_de_section(
        text["period"],
        "📅",
    )

    nam_hien_tai = date.today().year

    with st.container(border=True):
        col1, col2, col3 = st.columns(3)

        with col1:
            kieu = st.radio(
                text["type"],
                [
                    text["week"],
                    text["month"],
                ],
                horizontal=True,
            )

        with col2:
            nam = st.selectbox(
                text["year"],
                list(
                    range(
                        nam_hien_tai - 2,
                        nam_hien_tai + 2,
                    )
                ),
                index=2,
            )

        with col3:
            thang = st.selectbox(
                text["month_label"],
                list(range(1, 13)),
                index=date.today().month - 1,
            )

        danh_sach_bo_phan = (
            ThongKeController.lay_bo_phan()
        )

        bo_phan_options = [
            (None, text["all"])
        ] + [
            (
                int(row[0]),
                str(row[1]),
            )
            for row in danh_sach_bo_phan
        ]

        ten_bo_phan = st.selectbox(
            text["department"],
            [
                x[1]
                for x in bo_phan_options
            ],
        )

        bo_phan_id = next(
            x[0]
            for x in bo_phan_options
            if x[1] == ten_bo_phan
        )

        if kieu == text["week"]:
            cac_tuan = (
                _tao_cac_tuan_trong_thang(
                    nam,
                    thang,
                )
            )

            if not cac_tuan:
                st.warning(
                    "Không tạo được danh sách tuần."
                )
                return

            nhan_tuan = [
                (
                    f"Tuần {so}: "
                    f"{d1.strftime('%d/%m/%Y')} - "
                    f"{d2.strftime('%d/%m/%Y')}"
                )
                for so, d1, d2 in cac_tuan
            ]

            chon_tuan = st.selectbox(
                text["week_label"],
                nhan_tuan,
            )

            vi_tri = (
                nhan_tuan.index(
                    chon_tuan
                )
            )

            _, tu_ngay, den_ngay = (
                cac_tuan[vi_tri]
            )

            kieu_xuat = (
                f"Theo tuần - {chon_tuan}"
            )

        else:
            tu_ngay = date(
                nam,
                thang,
                1,
            )

            den_ngay = date(
                nam,
                thang,
                calendar.monthrange(
                    nam,
                    thang,
                )[1],
            )

            kieu_xuat = (
                f"Theo tháng - "
                f"{thang:02d}/{nam}"
            )

    hien_thi_tieu_de_section(
        text["overview"],
        "📌",
    )

    result = (
        ThongKeController.thong_ke_tong_hop(
            tu_ngay.strftime(
                "%Y-%m-%d"
            ),
            den_ngay.strftime(
                "%Y-%m-%d"
            ),
            bo_phan_id,
        )
    )

    if not result["success"]:
        st.error(
            result["message"]
        )
        return

    summary = result["summary"]
    rows = result["data"]

    c1, c2, c3, c4 = st.columns(4)

    with c1:
        hien_thi_the(
            text["people"],
            summary["tong_nguoi"],
            "👥",
        )

    with c2:
        hien_thi_the(
            text["days"],
            summary["tong_ngay_an"],
            "📅",
        )

    with c3:
        hien_thi_the(
            text["meals"],
            summary["tong_suat"],
            "🍚",
        )

    with c4:
        hien_thi_the(
            text["money"],
            f"{summary['tong_tien']:,.0f} đ",
            "💰",
        )

    hien_thi_tieu_de_section(
        text["detail"],
        "📋",
    )

    df = _tao_dataframe(rows)

    if df.empty:
        st.info(
            text["empty"]
        )
        return

    st.dataframe(
        df,
        width="stretch",
        hide_index=True,
        column_config={
            "STT": st.column_config.NumberColumn(
                "STT",
                format="%d",
            ),
            "Số ngày ăn": st.column_config.NumberColumn(
                "Số ngày ăn",
                format="%d",
            ),
            "Số suất": st.column_config.NumberColumn(
                "Số suất",
                format="%d",
            ),
            "Tiền cơm": st.column_config.NumberColumn(
                "Tiền cơm",
                format="%,.0f đ",
            ),
        },
    )

    excel_data = _xuat_excel(
        df,
        tu_ngay,
        den_ngay,
        kieu_xuat,
        ten_bo_phan,
        summary,
    )

    st.download_button(
        text["export"],
        data=excel_data,
        file_name=(
            "Bao_cao_thong_ke_suat_an_"
            f"{tu_ngay:%Y%m%d}_"
            f"{den_ngay:%Y%m%d}.xlsx"
        ),
        mime=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        use_container_width=True,
    )
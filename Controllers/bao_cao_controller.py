from Database.database import get_connection
from pathlib import Path
from datetime import date, timedelta
from copy import copy
import calendar
import csv

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


class BaoCaoController:

    @staticmethod
    def lay_du_lieu_bao_cao(tu_ngay, den_ngay, ten_nguoi=""):
        if not tu_ngay or not den_ngay:
            return {
                "success": False,
                "message": "Vui lòng chọn đầy đủ ngày.",
                "data": []
            }

        if tu_ngay > den_ngay:
            return {
                "success": False,
                "message": "Ngày bắt đầu không được lớn hơn ngày kết thúc.",
                "data": []
            }

        connection = get_connection()

        try:
            cursor = connection.cursor()

            cursor.execute("""
                SELECT a.Ngay,
                       n.HoTen,
                       n.SDT,
                       c.DaAn,
                       c.SoTienPhaiTra,
                       c.GhiChu
                FROM ChiTietAn c
                INNER JOIN NguoiAn n
                    ON c.NguoiAnId = n.Id
                INNER JOIN NgayAn a
                    ON c.NgayAnId = a.Id
                WHERE a.Ngay BETWEEN ? AND ?
                  AND n.HoTen LIKE ?
                ORDER BY a.Ngay, n.HoTen
            """, (
                tu_ngay,
                den_ngay,
                f"%{ten_nguoi.strip()}%"
            ))

            return {
                "success": True,
                "data": cursor.fetchall()
            }

        except Exception as error:
            return {
                "success": False,
                "message": f"Không thể lấy dữ liệu báo cáo: {error}",
                "data": []
            }

        finally:
            connection.close()

    @staticmethod
    def _format_ngay(ngay):
        if not ngay:
            return ""

        ngay = str(ngay)

        if len(ngay) >= 10:
            return f"{ngay[8:10]}/{ngay[5:7]}/{ngay[0:4]}"

        return ngay

    @staticmethod
    def _to_date(value):
        if isinstance(value, date):
            return value

        return date.fromisoformat(str(value)[:10])

    @staticmethod
    def _tuan_thu_hai(ngay):
        ngay = BaoCaoController._to_date(ngay)

        return ngay - timedelta(
            days=ngay.weekday()
        )

    @staticmethod
    def _lay_cac_tuan(tu_ngay, den_ngay):
        start_date = BaoCaoController._to_date(tu_ngay)
        end_date = BaoCaoController._to_date(den_ngay)

        monday = BaoCaoController._tuan_thu_hai(start_date)

        weeks = []

        while monday <= end_date:
            sunday = monday + timedelta(days=6)

            weeks.append({
                "tu_ngay": monday,
                "den_ngay": sunday
            })

            monday += timedelta(days=7)

        return weeks

    @staticmethod
    def _lay_cac_thang(tu_ngay, den_ngay):
        start_date = BaoCaoController._to_date(tu_ngay)
        end_date = BaoCaoController._to_date(den_ngay)

        current = date(
            start_date.year,
            start_date.month,
            1
        )

        months = []

        while current <= end_date:
            months.append({
                "nam": current.year,
                "thang": current.month,
                "tu_ngay": current,
                "den_ngay": date(
                    current.year,
                    current.month,
                    calendar.monthrange(
                        current.year,
                        current.month
                    )[1]
                )
            })

            if current.month == 12:
                current = date(
                    current.year + 1,
                    1,
                    1
                )
            else:
                current = date(
                    current.year,
                    current.month + 1,
                    1
                )

        return months

    @staticmethod
    def _thu_muc_report():
        report_dir = (
            Path(__file__).resolve().parent.parent
            / "Report"
        )

        report_dir.mkdir(
            parents=True,
            exist_ok=True
        )

        return report_dir

    @staticmethod
    def _duong_dan_report(file_path):
        report_dir = BaoCaoController._thu_muc_report()

        if file_path:
            file_name = Path(file_path).name
        else:
            file_name = "BaoCao.xlsx"

        original_path = report_dir / file_name

        if not original_path.exists():
            return original_path

        stem = original_path.stem
        suffix = original_path.suffix

        counter = 1

        while True:
            new_path = (
                report_dir
                / f"{stem}_{counter:02d}{suffix}"
            )

            if not new_path.exists():
                return new_path

            counter += 1

    @staticmethod
    def lay_bo_phan():
        connection = get_connection()

        try:
            cursor = connection.cursor()

            cursor.execute("""
                SELECT Id,
                       TenBoPhan
                FROM BoPhan
                WHERE DangHoatDong = 1
                ORDER BY TenBoPhan COLLATE NOCASE
            """)

            return cursor.fetchall()

        except Exception:
            return []

        finally:
            connection.close()

    @staticmethod
    def lay_danh_sach_nguoi_an(bo_phan_id=None):
        connection = get_connection()

        try:
            cursor = connection.cursor()

            sql = """
                SELECT n.Id,
                       n.HoTen,
                       COALESCE(
                           b.TenBoPhan,
                           'Chưa phân bộ phận'
                       ) AS TenBoPhan
                FROM NguoiAn n
                LEFT JOIN BoPhan b
                    ON n.BoPhanId = b.Id
            """

            params = []

            if bo_phan_id is not None:
                sql += """
                    WHERE n.BoPhanId = ?
                """

                params.append(bo_phan_id)

            sql += """
                ORDER BY n.HoTen COLLATE NOCASE
            """

            cursor.execute(
                sql,
                tuple(params)
            )

            return cursor.fetchall()

        finally:
            connection.close()

    @staticmethod
    def lay_nguoi_an_theo_id(nguoi_an_id):
        connection = get_connection()

        try:
            cursor = connection.cursor()

            cursor.execute("""
                SELECT n.Id,
                       n.HoTen,
                       n.SDT,
                       n.BoPhanId,
                       COALESCE(
                           b.TenBoPhan,
                           'Chưa phân bộ phận'
                       ) AS TenBoPhan,
                       n.DangHoatDong
                FROM NguoiAn n
                LEFT JOIN BoPhan b
                    ON n.BoPhanId = b.Id
                WHERE n.Id = ?
            """, (
                nguoi_an_id,
            ))

            return cursor.fetchone()

        finally:
            connection.close()

    @staticmethod
    def lay_ten_bo_phan(bo_phan_id):
        if bo_phan_id is None:
            return "Tất cả"

        connection = get_connection()

        try:
            cursor = connection.cursor()

            cursor.execute("""
                SELECT TenBoPhan
                FROM BoPhan
                WHERE Id = ?
            """, (
                bo_phan_id,
            ))

            row = cursor.fetchone()

            return (
                row[0]
                if row
                else "Không xác định"
            )

        finally:
            connection.close()

    @staticmethod
    def la_bo_phan_ve_sinh(ten_bo_phan):
        return (
            str(ten_bo_phan or "").strip().casefold()
            == "vệ sinh"
        )

    @staticmethod
    def _lay_du_lieu_tuan(
        tu_ngay,
        den_ngay,
        bo_phan_id=None
    ):
        connection = get_connection()

        try:
            cursor = connection.cursor()

            sql = """
                SELECT c.NguoiAnId,
                       a.Ngay,
                       c.DaAn,
                       c.SoTienPhaiTra
                FROM ChiTietAn c
                INNER JOIN NgayAn a
                    ON c.NgayAnId = a.Id
                INNER JOIN NguoiAn n
                    ON c.NguoiAnId = n.Id
                WHERE a.Ngay BETWEEN ? AND ?
                  AND c.DaAn = 1
            """

            params = [
                tu_ngay,
                den_ngay
            ]

            if bo_phan_id is not None:
                sql += """
                    AND n.BoPhanId = ?
                """

                params.append(bo_phan_id)

            cursor.execute(
                sql,
                tuple(params)
            )

            return cursor.fetchall()

        finally:
            connection.close()

    @staticmethod
    def _lay_giao_dich_tuan(
        nguoi_an_id,
        tu_ngay,
        den_ngay
    ):
        connection = get_connection()

        try:
            cursor = connection.cursor()

            cursor.execute("""
                SELECT COALESCE(
                    SUM(SoTien),
                    0
                )
                FROM GiaoDichNopTien
                WHERE NguoiAnId = ?
                  AND NgayNop BETWEEN ? AND ?
            """, (
                nguoi_an_id,
                tu_ngay,
                den_ngay
            ))

            row = cursor.fetchone()

            return float(
                row[0] or 0
            )

        finally:
            connection.close()

    @staticmethod
    def _lay_tien_phai_tra(
        nguoi_an_id,
        tu_ngay,
        den_ngay
    ):
        connection = get_connection()

        try:
            cursor = connection.cursor()

            cursor.execute("""
                SELECT COALESCE(
                    SUM(c.SoTienPhaiTra),
                    0
                )
                FROM ChiTietAn c
                INNER JOIN NgayAn a
                    ON c.NgayAnId = a.Id
                WHERE c.NguoiAnId = ?
                  AND c.DaAn = 1
                  AND a.Ngay BETWEEN ? AND ?
            """, (
                nguoi_an_id,
                tu_ngay,
                den_ngay
            ))

            row = cursor.fetchone()

            return float(
                row[0] or 0
            )

        finally:
            connection.close()

    @staticmethod
    def _lay_du_no_truoc_ky(
        nguoi_an_id,
        tu_ngay
    ):
        start = BaoCaoController._to_date(
            tu_ngay
        )

        previous_day = (
            start - timedelta(days=1)
        )

        connection = get_connection()

        try:
            cursor = connection.cursor()

            cursor.execute("""
                SELECT COALESCE(
                    SUM(SoTien),
                    0
                )
                FROM GiaoDichNopTien
                WHERE NguoiAnId = ?
                  AND NgayNop <= ?
            """, (
                nguoi_an_id,
                previous_day.isoformat()
            ))

            paid = float(
                cursor.fetchone()[0] or 0
            )

            cursor.execute("""
                SELECT COALESCE(
                    SUM(c.SoTienPhaiTra),
                    0
                )
                FROM ChiTietAn c
                INNER JOIN NgayAn a
                    ON c.NgayAnId = a.Id
                WHERE c.NguoiAnId = ?
                  AND c.DaAn = 1
                  AND a.Ngay <= ?
            """, (
                nguoi_an_id,
                previous_day.isoformat()
            ))

            due = float(
                cursor.fetchone()[0] or 0
            )

            return paid - due

        finally:
            connection.close()

    @staticmethod
    def _lay_du_no_luy_ke(
        nguoi_an_id,
        den_ngay
    ):
        connection = get_connection()

        try:
            cursor = connection.cursor()

            cursor.execute("""
                SELECT COALESCE(
                    SUM(SoTien),
                    0
                )
                FROM GiaoDichNopTien
                WHERE NguoiAnId = ?
                  AND NgayNop <= ?
            """, (
                nguoi_an_id,
                den_ngay
            ))

            paid = float(
                cursor.fetchone()[0] or 0
            )

            cursor.execute("""
                SELECT COALESCE(
                    SUM(c.SoTienPhaiTra),
                    0
                )
                FROM ChiTietAn c
                INNER JOIN NgayAn a
                    ON c.NgayAnId = a.Id
                WHERE c.NguoiAnId = ?
                  AND c.DaAn = 1
                  AND a.Ngay <= ?
            """, (
                nguoi_an_id,
                den_ngay
            ))

            due = float(
                cursor.fetchone()[0] or 0
            )

            return paid - due

        finally:
            connection.close()

    @staticmethod
    def _lay_thong_tin_thanh_toan(
        nguoi_an_id,
        tu_ngay,
        den_ngay
    ):
        start = BaoCaoController._to_date(
            tu_ngay
        )

        end = BaoCaoController._to_date(
            den_ngay
        )

        opening = BaoCaoController._lay_du_no_truoc_ky(
            nguoi_an_id,
            start
        )

        period_due = BaoCaoController._lay_tien_phai_tra(
            nguoi_an_id,
            start.isoformat(),
            end.isoformat()
        )

        period_paid = BaoCaoController._lay_giao_dich_tuan(
            nguoi_an_id,
            start.isoformat(),
            end.isoformat()
        )

        period_balance = (
            period_paid - period_due
        )

        closing = (
            opening + period_balance
        )

        return {
            "NoDauKy": opening,
            "PhaiTra": period_due,
            "DaNop": period_paid,
            "SoDuKy": period_balance,
            "SoDuCuoiKy": closing,

            "ConNoKy": (
                abs(period_balance)
                if period_balance < -0.01
                else 0
            ),

            "NopThuaKy": (
                period_balance
                if period_balance > 0.01
                else 0
            ),

            "ConNo": (
                abs(closing)
                if closing < -0.01
                else 0
            ),

            "NopThua": (
                closing
                if closing > 0.01
                else 0
            ),

            "TrangThaiKy": (
                "Còn nợ"
                if period_balance < -0.01
                else "Đã đủ"
            ),

            "TrangThai": (
                "Còn nợ"
                if closing < -0.01
                else "Đã đủ"
            )
        }

    @staticmethod
    def _lay_thong_tin_an(
        nguoi_an_id,
        tu_ngay,
        den_ngay
    ):
        connection = get_connection()

        try:
            cursor = connection.cursor()

            cursor.execute("""
                SELECT COUNT(*),
                       COALESCE(
                           SUM(
                               c.SoTienPhaiTra
                           ),
                           0
                       )
                FROM ChiTietAn c
                INNER JOIN NgayAn a
                    ON c.NgayAnId = a.Id
                WHERE c.NguoiAnId = ?
                  AND c.DaAn = 1
                  AND a.Ngay BETWEEN ? AND ?
            """, (
                nguoi_an_id,
                tu_ngay,
                den_ngay
            ))

            row = cursor.fetchone()

            total_meals = int(
                row[0] or 0
            )

            total_due = float(
                row[1] or 0
            )

            return {
                "TongBua": total_meals,
                "ThanhTien": total_due,
                "DonGia": (
                    total_due / total_meals
                    if total_meals
                    else 0
                )
            }

        finally:
            connection.close()

    @staticmethod
    def _tao_dong_thanh_toan(
        nguoi_id,
        ho_ten,
        bo_phan,
        tu_ngay,
        den_ngay,
        thong_tin_an=None
    ):
        if thong_tin_an is None:
            thong_tin_an = (
                BaoCaoController
                ._lay_thong_tin_an(
                    nguoi_id,
                    tu_ngay,
                    den_ngay
                )
            )

        payment = (
            BaoCaoController
            ._lay_thong_tin_thanh_toan(
                nguoi_id,
                tu_ngay,
                den_ngay
            )
        )

        return {
            "NguoiId": nguoi_id,
            "HoTen": ho_ten,
            "BoPhan": bo_phan,
            "TuNgay": BaoCaoController._to_date(
                tu_ngay
            ),
            "DenNgay": BaoCaoController._to_date(
                den_ngay
            ),

            "TongBua": thong_tin_an[
                "TongBua"
            ],

            "DonGia": thong_tin_an[
                "DonGia"
            ],

            "ThanhTien": thong_tin_an[
                "ThanhTien"
            ],

            "NoDauKy": payment[
                "NoDauKy"
            ],

            "PhaiTra": payment[
                "PhaiTra"
            ],

            "DaDong": payment[
                "DaNop"
            ],

            "DaNop": payment[
                "DaNop"
            ],

            "SoDuKy": payment[
                "SoDuKy"
            ],

            "SoDuCuoiKy": payment[
                "SoDuCuoiKy"
            ],

            "ConNoKy": payment[
                "ConNoKy"
            ],

            "NopThuaKy": payment[
                "NopThuaKy"
            ],

            "ConNo": payment[
                "ConNo"
            ],

            "NopThua": payment[
                "NopThua"
            ],

            "TrangThaiKy": payment[
                "TrangThaiKy"
            ],

            "TrangThai": payment[
                "TrangThai"
            ]
        }

    @staticmethod
    def lay_bao_cao_tuan(
        tu_ngay,
        bo_phan_id=None
    ):
        try:
            start = BaoCaoController._tuan_thu_hai(
                tu_ngay
            )

            days = [
                start + timedelta(days=i)
                for i in range(7)
            ]

            end = days[-1]

            people = (
                BaoCaoController
                .lay_danh_sach_nguoi_an(
                    bo_phan_id
                )
            )

            meals = (
                BaoCaoController
                ._lay_du_lieu_tuan(
                    start.isoformat(),
                    end.isoformat(),
                    bo_phan_id
                )
            )

            meal_map = {}
            price_map = {}

            for (
                nguoi_id,
                ngay,
                da_an,
                so_tien
            ) in meals:

                if da_an != 1:
                    continue

                ngay_key = str(ngay)[:10]

                meal_map[
                    (
                        nguoi_id,
                        ngay_key
                    )
                ] = True

                price_map.setdefault(
                    nguoi_id,
                    []
                ).append(
                    float(
                        so_tien or 0
                    )
                )

            rows = []

            for (
                nguoi_id,
                ho_ten,
                bo_phan
            ) in people:

                marks = []

                for day in days:
                    marks.append(
                        "☑"
                        if meal_map.get(
                            (
                                nguoi_id,
                                day.isoformat()
                            ),
                            False
                        )
                        else "☐"
                    )

                thong_tin_an = {
                    "TongBua": sum(
                        1
                        for mark in marks
                        if mark == "☑"
                    ),
                    "ThanhTien": sum(
                        price_map.get(
                            nguoi_id,
                            []
                        )
                    )
                }

                thong_tin_an["DonGia"] = (
                    thong_tin_an["ThanhTien"]
                    / thong_tin_an["TongBua"]
                    if thong_tin_an["TongBua"]
                    else 0
                )

                row = (
                    BaoCaoController
                    ._tao_dong_thanh_toan(
                        nguoi_id,
                        ho_ten,
                        bo_phan,
                        start,
                        end,
                        thong_tin_an
                    )
                )

                row["Marks"] = marks

                rows.append(row)

            return {
                "success": True,
                "tu_ngay": start,
                "den_ngay": end,
                "data": rows
            }

        except Exception as error:
            return {
                "success": False,
                "message": (
                    f"Không thể tạo báo cáo tuần: {error}"
                ),
                "data": []
            }

    @staticmethod
    def lay_bao_cao_thang(
        nam,
        thang,
        bo_phan_id=None
    ):
        try:
            nam = int(nam)
            thang = int(thang)

            start = date(
                nam,
                thang,
                1
            )

            end = date(
                nam,
                thang,
                calendar.monthrange(
                    nam,
                    thang
                )[1]
            )

            people = (
                BaoCaoController
                .lay_danh_sach_nguoi_an(
                    bo_phan_id
                )
            )

            rows = []

            for (
                nguoi_id,
                ho_ten,
                bo_phan
            ) in people:

                row = (
                    BaoCaoController
                    ._tao_dong_thanh_toan(
                        nguoi_id,
                        ho_ten,
                        bo_phan,
                        start,
                        end
                    )
                )

                rows.append(row)

            return {
                "success": True,
                "tu_ngay": start,
                "den_ngay": end,
                "data": rows
            }

        except Exception as error:
            return {
                "success": False,
                "message": (
                    f"Không thể tạo báo cáo tháng: {error}"
                ),
                "data": []
            }

    @staticmethod
    def lay_tong_quan_thanh_toan(
        tu_ngay,
        den_ngay,
        bo_phan_id=None
    ):
        try:
            start = BaoCaoController._to_date(
                tu_ngay
            )

            end = BaoCaoController._to_date(
                den_ngay
            )

            if start > end:
                return {
                    "success": False,
                    "message": (
                        "Ngày bắt đầu không được lớn hơn "
                        "ngày kết thúc."
                    ),
                    "data": [],
                    "tong_quan": {}
                }

            people = (
                BaoCaoController
                .lay_danh_sach_nguoi_an(
                    bo_phan_id
                )
            )

            rows = []

            for (
                nguoi_id,
                ho_ten,
                bo_phan
            ) in people:

                row = (
                    BaoCaoController
                    ._tao_dong_thanh_toan(
                        nguoi_id,
                        ho_ten,
                        bo_phan,
                        start,
                        end
                    )
                )

                rows.append(row)

            tong_phai_tra = sum(
                row["PhaiTra"]
                for row in rows
            )

            tong_da_nop = sum(
                row["DaNop"]
                for row in rows
            )

            tong_con_no = sum(
                row["ConNo"]
                for row in rows
            )

            tong_nop_thua = sum(
                row["NopThua"]
                for row in rows
            )

            so_da_du = sum(
                1
                for row in rows
                if row["TrangThai"] == "Đã đủ"
            )

            so_con_no = sum(
                1
                for row in rows
                if row["TrangThai"] == "Còn nợ"
            )

            so_nop_thua = sum(
                1
                for row in rows
                if row["NopThua"] > 0.01
            )

            return {
                "success": True,
                "tu_ngay": start,
                "den_ngay": end,
                "data": rows,
                "tong_quan": {
                    "TongPhaiTra": tong_phai_tra,
                    "TongDaNop": tong_da_nop,
                    "TongConNo": tong_con_no,
                    "TongNopThua": tong_nop_thua,
                    "SoDaDu": so_da_du,
                    "SoConNo": so_con_no,
                    "SoNopThua": so_nop_thua,
                    "TongNguoi": len(rows)
                }
            }

        except Exception as error:
            return {
                "success": False,
                "message": (
                    f"Không thể tạo tổng quan thanh toán: {error}"
                ),
                "data": [],
                "tong_quan": {}
            }

    @staticmethod
    def lay_chi_tiet_con_no(
        tu_ngay,
        den_ngay,
        bo_phan_id=None
    ):
        result = (
            BaoCaoController
            .lay_tong_quan_thanh_toan(
                tu_ngay,
                den_ngay,
                bo_phan_id
            )
        )

        if not result["success"]:
            return result

        rows = [
            row
            for row in result["data"]
            if row["ConNo"] > 0.01
        ]

        rows.sort(
            key=lambda item: item["ConNo"],
            reverse=True
        )

        return {
            "success": True,
            "tu_ngay": result["tu_ngay"],
            "den_ngay": result["den_ngay"],
            "data": rows,
            "tong_tien": sum(
                row["ConNo"]
                for row in rows
            )
        }

    @staticmethod
    def lay_chi_tiet_nop_thua(
        tu_ngay,
        den_ngay,
        bo_phan_id=None
    ):
        result = (
            BaoCaoController
            .lay_tong_quan_thanh_toan(
                tu_ngay,
                den_ngay,
                bo_phan_id
            )
        )

        if not result["success"]:
            return result

        rows = [
            row
            for row in result["data"]
            if row["NopThua"] > 0.01
        ]

        rows.sort(
            key=lambda item: item["NopThua"],
            reverse=True
        )

        return {
            "success": True,
            "tu_ngay": result["tu_ngay"],
            "den_ngay": result["den_ngay"],
            "data": rows,
            "tong_tien": sum(
                row["NopThua"]
                for row in rows
            )
        }

    @staticmethod
    def lay_chi_tiet_da_du(
        tu_ngay,
        den_ngay,
        bo_phan_id=None
    ):
        result = (
            BaoCaoController
            .lay_tong_quan_thanh_toan(
                tu_ngay,
                den_ngay,
                bo_phan_id
            )
        )

        if not result["success"]:
            return result

        rows = [
            row
            for row in result["data"]
            if row["TrangThai"] == "Đã đủ"
        ]

        rows.sort(
            key=lambda item: item["HoTen"].casefold()
        )

        return {
            "success": True,
            "tu_ngay": result["tu_ngay"],
            "den_ngay": result["den_ngay"],
            "data": rows,
            "so_nguoi": len(rows)
        }

    @staticmethod
    def lay_chi_tiet_phai_tra(
        tu_ngay,
        den_ngay,
        bo_phan_id=None
    ):
        result = (
            BaoCaoController
            .lay_tong_quan_thanh_toan(
                tu_ngay,
                den_ngay,
                bo_phan_id
            )
        )

        if not result["success"]:
            return result

        return {
            "success": True,
            "tu_ngay": result["tu_ngay"],
            "den_ngay": result["den_ngay"],
            "data": result["data"],
            "tong_tien": (
                result["tong_quan"]
                ["TongPhaiTra"]
            )
        }

    @staticmethod
    def lay_chi_tiet_da_nop(
        tu_ngay,
        den_ngay,
        bo_phan_id=None
    ):
        result = (
            BaoCaoController
            .lay_tong_quan_thanh_toan(
                tu_ngay,
                den_ngay,
                bo_phan_id
            )
        )

        if not result["success"]:
            return result

        return {
            "success": True,
            "tu_ngay": result["tu_ngay"],
            "den_ngay": result["den_ngay"],
            "data": result["data"],
            "tong_tien": (
                result["tong_quan"]
                ["TongDaNop"]
            )
        }

    @staticmethod
    def lay_bao_cao_tuan_theo_khoang(
        tu_ngay,
        den_ngay,
        bo_phan_id=None
    ):
        try:
            start = BaoCaoController._to_date(
                tu_ngay
            )

            end = BaoCaoController._to_date(
                den_ngay
            )

            weeks = (
                BaoCaoController
                ._lay_cac_tuan(
                    start,
                    end
                )
            )

            data = []

            for index, week in enumerate(
                weeks,
                1
            ):
                result = (
                    BaoCaoController
                    .lay_bao_cao_tuan(
                        week["tu_ngay"],
                        bo_phan_id
                    )
                )

                if not result["success"]:
                    continue

                rows = result["data"]

                tong_phai_tra = sum(
                    row["PhaiTra"]
                    for row in rows
                )

                tong_da_nop = sum(
                    row["DaNop"]
                    for row in rows
                )

                tong_con_no = sum(
                    row["ConNo"]
                    for row in rows
                )

                tong_nop_thua = sum(
                    row["NopThua"]
                    for row in rows
                )

                tong_con_no_ky = sum(
                    row["ConNoKy"]
                    for row in rows
                )

                tong_nop_thua_ky = sum(
                    row["NopThuaKy"]
                    for row in rows
                )

                so_da_du = sum(
                    1
                    for row in rows
                    if row["TrangThaiKy"] == "Đã đủ"
                )

                so_con_no = sum(
                    1
                    for row in rows
                    if row["TrangThaiKy"] == "Còn nợ"
                )

                so_nop_thua = sum(
                    1
                    for row in rows
                    if row["NopThuaKy"] > 0.01
                )

                data.append({
                    "Tuan": index,
                    "TuNgay": week["tu_ngay"],
                    "DenNgay": week["den_ngay"],

                    "PhaiTra": tong_phai_tra,
                    "DaNop": tong_da_nop,

                    "ConNo": tong_con_no,
                    "NopThua": tong_nop_thua,

                    "ConNoKy": tong_con_no_ky,
                    "NopThuaKy": tong_nop_thua_ky,

                    "SoDaDu": so_da_du,
                    "SoConNo": so_con_no,
                    "SoNopThua": so_nop_thua,

                    "DaQuyetToanTuan": (
                        so_con_no == 0
                    ),

                    "TrangThai": (
                        "Còn nợ"
                        if so_con_no > 0
                        else "Đã đủ"
                    ),

                    "TrangThaiKy": (
                        "Còn nợ"
                        if so_con_no > 0
                        else "Đã đủ"
                    ),

                    "ChiTiet": rows
                })

            return {
                "success": True,
                "data": data
            }

        except Exception as error:
            return {
                "success": False,
                "message": (
                    f"Không thể tạo tổng hợp tuần: {error}"
                ),
                "data": []
            }

    @staticmethod
    def lay_bao_cao_thang_theo_khoang(
        tu_ngay,
        den_ngay,
        bo_phan_id=None
    ):
        try:
            start = BaoCaoController._to_date(
                tu_ngay
            )

            end = BaoCaoController._to_date(
                den_ngay
            )

            months = (
                BaoCaoController
                ._lay_cac_thang(
                    start,
                    end
                )
            )

            data = []

            for item in months:
                result = (
                    BaoCaoController
                    .lay_bao_cao_thang(
                        item["nam"],
                        item["thang"],
                        bo_phan_id
                    )
                )

                if not result["success"]:
                    continue

                rows = result["data"]

                tong_phai_tra = sum(
                    row["PhaiTra"]
                    for row in rows
                )

                tong_da_nop = sum(
                    row["DaNop"]
                    for row in rows
                )

                tong_con_no = sum(
                    row["ConNo"]
                    for row in rows
                )

                tong_nop_thua = sum(
                    row["NopThua"]
                    for row in rows
                )

                tong_con_no_ky = sum(
                    row["ConNoKy"]
                    for row in rows
                )

                tong_nop_thua_ky = sum(
                    row["NopThuaKy"]
                    for row in rows
                )

                so_da_du = sum(
                    1
                    for row in rows
                    if row["TrangThaiKy"] == "Đã đủ"
                )

                so_con_no = sum(
                    1
                    for row in rows
                    if row["TrangThaiKy"] == "Còn nợ"
                )

                so_nop_thua = sum(
                    1
                    for row in rows
                    if row["NopThuaKy"] > 0.01
                )

                data.append({
                    "Nam": item["nam"],
                    "Thang": item["thang"],
                    "TuNgay": item["tu_ngay"],
                    "DenNgay": item["den_ngay"],

                    "PhaiTra": tong_phai_tra,
                    "DaNop": tong_da_nop,

                    "ConNo": tong_con_no,
                    "NopThua": tong_nop_thua,

                    "ConNoKy": tong_con_no_ky,
                    "NopThuaKy": tong_nop_thua_ky,

                    "SoDaDu": so_da_du,
                    "SoConNo": so_con_no,
                    "SoNopThua": so_nop_thua,

                    "DaQuyetToanThang": (
                        so_con_no == 0
                    ),

                    "TrangThai": (
                        "Còn nợ"
                        if so_con_no > 0
                        else "Đã đủ"
                    ),

                    "TrangThaiKy": (
                        "Còn nợ"
                        if so_con_no > 0
                        else "Đã đủ"
                    ),

                    "ChiTiet": rows
                })

            return {
                "success": True,
                "data": data
            }

        except Exception as error:
            return {
                "success": False,
                "message": (
                    f"Không thể tạo tổng hợp tháng: {error}"
                ),
                "data": []
            }

    @staticmethod
    def lay_quyet_toan_ve_sinh_theo_thang(
        nam,
        thang
    ):
        try:
            nam = int(nam)
            thang = int(thang)

            bo_phan_rows = (
                BaoCaoController
                .lay_bo_phan()
            )

            ve_sinh_id = None
            ten_ve_sinh = None

            for row in bo_phan_rows:
                if BaoCaoController.la_bo_phan_ve_sinh(
                    row[1]
                ):
                    ve_sinh_id = row[0]
                    ten_ve_sinh = row[1]
                    break

            if ve_sinh_id is None:
                return {
                    "success": True,
                    "message": (
                        "Không tìm thấy bộ phận Vệ sinh."
                    ),
                    "bo_phan_id": None,
                    "bo_phan": "Vệ sinh",
                    "data": [],
                    "tong_quan": {}
                }

            result = (
                BaoCaoController
                .lay_bao_cao_thang(
                    nam,
                    thang,
                    ve_sinh_id
                )
            )

            if not result["success"]:
                return result

            rows = result["data"]

            tong_quan = {
                "TongPhaiTra": sum(
                    row["PhaiTra"]
                    for row in rows
                ),

                "TongDaNop": sum(
                    row["DaNop"]
                    for row in rows
                ),

                "TongConNo": sum(
                    row["ConNo"]
                    for row in rows
                ),

                "TongNopThua": sum(
                    row["NopThua"]
                    for row in rows
                ),

                "SoDaDu": sum(
                    1
                    for row in rows
                    if row["TrangThai"] == "Đã đủ"
                ),

                "SoConNo": sum(
                    1
                    for row in rows
                    if row["TrangThai"] == "Còn nợ"
                ),

                "SoNopThua": sum(
                    1
                    for row in rows
                    if row["NopThua"] > 0.01
                ),

                "TongNguoi": len(rows)
            }

            return {
                "success": True,
                "bo_phan_id": ve_sinh_id,
                "bo_phan": (
                    ten_ve_sinh
                    or "Vệ sinh"
                ),
                "tu_ngay": result["tu_ngay"],
                "den_ngay": result["den_ngay"],
                "data": rows,
                "tong_quan": tong_quan
            }

        except Exception as error:
            return {
                "success": False,
                "message": (
                    "Không thể quyết toán "
                    f"bộ phận Vệ sinh: {error}"
                ),
                "data": [],
                "tong_quan": {}
            }

    @staticmethod
    def lay_theo_doi_ve_sinh_theo_tuan(
        tu_ngay,
        den_ngay
    ):
        try:
            bo_phan_rows = (
                BaoCaoController
                .lay_bo_phan()
            )

            ve_sinh_id = None

            for row in bo_phan_rows:
                if BaoCaoController.la_bo_phan_ve_sinh(
                    row[1]
                ):
                    ve_sinh_id = row[0]
                    break

            if ve_sinh_id is None:
                return {
                    "success": True,
                    "data": []
                }

            return (
                BaoCaoController
                .lay_bao_cao_tuan_theo_khoang(
                    tu_ngay,
                    den_ngay,
                    ve_sinh_id
                )
            )

        except Exception as error:
            return {
                "success": False,
                "message": (
                    "Không thể theo dõi "
                    f"Vệ sinh theo tuần: {error}"
                ),
                "data": []
            }

    @staticmethod
    def lay_tong_quan_theo_bo_phan(
        tu_ngay,
        den_ngay
    ):
        try:
            departments = (
                BaoCaoController
                .lay_bo_phan()
            )

            data = []

            for (
                bo_phan_id,
                ten_bo_phan
            ) in departments:

                result = (
                    BaoCaoController
                    .lay_tong_quan_thanh_toan(
                        tu_ngay,
                        den_ngay,
                        bo_phan_id
                    )
                )

                if not result["success"]:
                    continue

                summary = result[
                    "tong_quan"
                ]

                data.append({
                    "BoPhanId": bo_phan_id,
                    "BoPhan": ten_bo_phan,
                    "TongNguoi": summary[
                        "TongNguoi"
                    ],
                    "PhaiTra": summary[
                        "TongPhaiTra"
                    ],
                    "DaNop": summary[
                        "TongDaNop"
                    ],
                    "ConNo": summary[
                        "TongConNo"
                    ],
                    "NopThua": summary[
                        "TongNopThua"
                    ],
                    "SoDaDu": summary[
                        "SoDaDu"
                    ],
                    "SoConNo": summary[
                        "SoConNo"
                    ],
                    "SoNopThua": summary[
                        "SoNopThua"
                    ],
                    "HinhThucQuyetToan": (
                        "Tháng"
                        if BaoCaoController
                        .la_bo_phan_ve_sinh(
                            ten_bo_phan
                        )
                        else "Tuần"
                    )
                })

            return {
                "success": True,
                "data": data
            }

        except Exception as error:
            return {
                "success": False,
                "message": (
                    "Không thể tổng hợp "
                    f"theo bộ phận: {error}"
                ),
                "data": []
            }

    @staticmethod
    def _tao_style_excel():
        font_normal = Font(
            name="Times New Roman",
            size=13
        )

        font_bold = Font(
            name="Times New Roman",
            size=13,
            bold=True
        )

        font_title = Font(
            name="Times New Roman",
            size=13,
            bold=True
        )

        border = Border(
            left=Side(
                style="thin",
                color="BFBFBF"
            ),
            right=Side(
                style="thin",
                color="BFBFBF"
            ),
            top=Side(
                style="thin",
                color="BFBFBF"
            ),
            bottom=Side(
                style="thin",
                color="BFBFBF"
            )
        )

        title_fill = PatternFill(
            "solid",
            fgColor="E8DDEB"
        )

        header_fill = PatternFill(
            "solid",
            fgColor="D9EAF7"
        )

        total_fill = PatternFill(
            "solid",
            fgColor="F3E5F5"
        )

        checked_fill = PatternFill(
            "solid",
            fgColor="E2F0D9"
        )

        unchecked_fill = PatternFill(
            "solid",
            fgColor="F2F2F2"
        )

        # Zebra row: hai màu trung tính, sang và có độ tương phản vừa đủ
        # để mắt dễ bám theo từng dòng khi đối chiếu dữ liệu.
        # Không dùng nhiều màu khác nhau giữa từng người vì sẽ làm bảng rối.
        row_fills = [
            PatternFill("solid", fgColor="F8FAFC"),  # White Smoke
            PatternFill("solid", fgColor="EAF0F6"),  # Soft Slate Blue
        ]

        # Dòng còn nợ: đỏ nhạt, rõ nhưng không làm mất khả năng đọc.
        debt_fill = PatternFill(
            "solid",
            fgColor="FDE2E1"
        )

        center = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        left = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True
        )

        right = Alignment(
            horizontal="right",
            vertical="center"
        )

        money_format = '#,##0 "đ"'

        return {
            "font_normal": font_normal,
            "font_bold": font_bold,
            "font_title": font_title,
            "border": border,
            "title_fill": title_fill,
            "header_fill": header_fill,
            "total_fill": total_fill,
            "checked_fill": checked_fill,
            "unchecked_fill": unchecked_fill,
            "row_fills": row_fills,
            "debt_fill": debt_fill,
            "center": center,
            "left": left,
            "right": right,
            "money_format": money_format
        }

    @staticmethod
    def _ap_dung_mau_nguoi(
        sheet,
        row,
        start_col,
        end_col,
        ho_ten,
        con_no,
        styles
    ):
        """Tô màu một dòng theo từng người; còn nợ thì tô đỏ nhạt."""
        try:
            debt = float(con_no or 0) > 0.01
        except (TypeError, ValueError):
            debt = False

        if debt:
            fill = styles["debt_fill"]
        else:
            # Xen kẽ theo đúng từng dòng dữ liệu.
            # Dùng số dòng thay vì tên người để khi nhìn ngang,
            # mỗi dòng được tách bạch rõ ràng và không bị loè màu.
            fills = styles["row_fills"]
            fill = fills[(row - 1) % len(fills)]

        for col in range(start_col, end_col + 1):
            sheet.cell(row, col).fill = fill

    @staticmethod
    def _set_excel_title(
        sheet,
        row,
        title,
        end_col,
        styles
    ):
        sheet.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=end_col
        )

        cell = sheet.cell(
            row,
            1
        )

        cell.value = title
        cell.font = styles["font_title"]
        cell.fill = styles["title_fill"]
        cell.alignment = styles["center"]
        cell.border = styles["border"]

        sheet.row_dimensions[row].height = 25

    @staticmethod
    def _set_excel_header(
        sheet,
        row,
        headers,
        styles
    ):
        for col, header in enumerate(
            headers,
            1
        ):
            cell = sheet.cell(
                row,
                col
            )

            cell.value = header
            cell.font = styles["font_bold"]
            cell.fill = styles["header_fill"]
            cell.alignment = styles["center"]
            cell.border = styles["border"]

        sheet.row_dimensions[row].height = 30

    @staticmethod
    def _set_excel_total(
        sheet,
        row,
        start_col,
        end_col,
        styles
    ):
        for col in range(
            start_col,
            end_col + 1
        ):
            cell = sheet.cell(
                row,
                col
            )

            cell.font = styles["font_bold"]
            cell.fill = styles["total_fill"]
            cell.border = styles["border"]
            cell.alignment = styles["center"]

    @staticmethod
    def _set_excel_widths(
        sheet,
        widths
    ):
        for col, width in enumerate(
            widths,
            1
        ):
            sheet.column_dimensions[
                get_column_letter(col)
            ].width = width

    @staticmethod
    def _finalize_excel_sheet(sheet):
        sheet.sheet_view.showGridLines = False

        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue

                bold = bool(
                    cell.font.bold
                    if cell.font
                    else False
                )

                italic = bool(
                    cell.font.italic
                    if cell.font
                    else False
                )

                cell.font = Font(
                    name="Times New Roman",
                    size=13,
                    bold=bold,
                    italic=italic
                )

                if cell.border:
                    cell.border = copy(
                        cell.border
                    )

        sheet.page_setup.orientation = "landscape"

        sheet.page_margins.left = 0.25
        sheet.page_margins.right = 0.25
        sheet.page_margins.top = 0.5
        sheet.page_margins.bottom = 0.5

        sheet.sheet_properties.pageSetUpPr.fitToPage = False

    @staticmethod
    def _tao_sheet_bao_cao(
        workbook,
        du_lieu,
        styles
    ):
        sheet = workbook.active
        sheet.title = "Báo cáo"

        du_lieu_theo_ngay = {}

        for row in du_lieu:
            du_lieu_theo_ngay.setdefault(
                row[0],
                []
            ).append(row)

        current_row = 1

        for ngay in sorted(
            du_lieu_theo_ngay
        ):
            BaoCaoController._set_excel_title(
                sheet,
                current_row,
                (
                    "CHẤM CƠM NGÀY "
                    f"{BaoCaoController._format_ngay(ngay)}"
                ),
                7,
                styles
            )

            current_row += 2

            headers = [
                "STT",
                "Ngày",
                "Họ tên",
                "SĐT",
                "Đã ăn",
                "Số tiền phải trả",
                "Ghi chú"
            ]

            BaoCaoController._set_excel_header(
                sheet,
                current_row,
                headers,
                styles
            )

            current_row += 1

            danh_sach = du_lieu_theo_ngay[
                ngay
            ]

            for stt, row in enumerate(
                danh_sach,
                1
            ):
                values = [
                    stt,
                    BaoCaoController._format_ngay(
                        row[0]
                    ),
                    row[1],
                    row[2] or "",
                    "Có"
                    if row[3] == 1
                    else "Không",
                    row[4] or 0,
                    row[5] or ""
                ]

                for col, value in enumerate(
                    values,
                    1
                ):
                    cell = sheet.cell(
                        current_row,
                        col,
                        value
                    )

                    cell.font = styles[
                        "font_normal"
                    ]

                    cell.border = styles[
                        "border"
                    ]

                    if col in [
                        1,
                        2,
                        4,
                        5
                    ]:
                        cell.alignment = (
                            styles["center"]
                        )
                    elif col == 6:
                        cell.alignment = (
                            styles["right"]
                        )

                        cell.number_format = (
                            styles["money_format"]
                        )
                    else:
                        cell.alignment = (
                            styles["left"]
                        )

                BaoCaoController._ap_dung_mau_nguoi(
                    sheet,
                    current_row,
                    1,
                    len(headers),
                    row[1],
                    0,
                    styles
                )

                current_row += 1

            tong_luot_an = sum(
                1
                for row in danh_sach
                if row[3] == 1
            )

            tong_tien_ngay = sum(
                (row[4] or 0)
                for row in danh_sach
                if row[3] == 1
            )

            sheet.merge_cells(
                start_row=current_row,
                start_column=1,
                end_row=current_row,
                end_column=5
            )

            cell = sheet.cell(
                current_row,
                1
            )

            cell.value = (
                f"Tổng lượt ăn: {tong_luot_an}"
            )

            cell.font = styles[
                "font_bold"
            ]

            cell.fill = styles[
                "total_fill"
            ]

            cell.alignment = styles[
                "right"
            ]

            cell.border = styles[
                "border"
            ]

            cell = sheet.cell(
                current_row,
                6
            )

            cell.value = tong_tien_ngay

            cell.font = styles[
                "font_bold"
            ]

            cell.fill = styles[
                "total_fill"
            ]

            cell.alignment = styles[
                "right"
            ]

            cell.number_format = styles[
                "money_format"
            ]

            cell.border = styles[
                "border"
            ]

            sheet.cell(
                current_row,
                7
            ).fill = styles[
                "total_fill"
            ]

            sheet.cell(
                current_row,
                7
            ).border = styles[
                "border"
            ]

            current_row += 2

        BaoCaoController._set_excel_widths(
            sheet,
            [
                8,
                13,
                28,
                16,
                12,
                20,
                30
            ]
        )

        sheet.freeze_panes = "A3"

        return sheet

    @staticmethod
    def _tao_sheet_theo_ngay(
        workbook,
        du_lieu,
        styles
    ):
        sheet = workbook.create_sheet(
            "Theo ngày"
        )

        headers = [
            "Ngày",
            "Số người ăn",
            "Tổng tiền"
        ]

        BaoCaoController._set_excel_header(
            sheet,
            1,
            headers,
            styles
        )

        thong_ke = {}

        for row in du_lieu:
            if row[0] not in thong_ke:
                thong_ke[row[0]] = [
                    0,
                    0
                ]

            if row[3] == 1:
                thong_ke[row[0]][0] += 1
                thong_ke[row[0]][1] += (
                    row[4] or 0
                )

        current_row = 2

        for ngay in sorted(thong_ke):
            values = [
                BaoCaoController._format_ngay(
                    ngay
                ),
                thong_ke[ngay][0],
                thong_ke[ngay][1]
            ]

            for col, value in enumerate(
                values,
                1
            ):
                cell = sheet.cell(
                    current_row,
                    col,
                    value
                )

                cell.font = styles[
                    "font_normal"
                ]

                cell.border = styles[
                    "border"
                ]

                if col == 3:
                    cell.number_format = (
                        styles["money_format"]
                    )

                    cell.alignment = (
                        styles["right"]
                    )
                else:
                    cell.alignment = (
                        styles["center"]
                    )

            current_row += 1

        BaoCaoController._set_excel_widths(
            sheet,
            [
                18,
                18,
                22
            ]
        )

        sheet.freeze_panes = "A2"

        if current_row > 2:
            sheet.auto_filter.ref = (
                f"A1:C{current_row - 1}"
            )

        return sheet

    @staticmethod
    def _tao_sheet_theo_nguoi(
        workbook,
        du_lieu,
        styles
    ):
        sheet = workbook.create_sheet(
            "Theo người"
        )

        headers = [
            "Họ tên",
            "Số lần ăn",
            "Tổng tiền"
        ]

        BaoCaoController._set_excel_header(
            sheet,
            1,
            headers,
            styles
        )

        thong_ke = {}

        for row in du_lieu:
            ho_ten = row[1]

            if ho_ten not in thong_ke:
                thong_ke[ho_ten] = [
                    0,
                    0
                ]

            if row[3] == 1:
                thong_ke[ho_ten][0] += 1
                thong_ke[ho_ten][1] += (
                    row[4] or 0
                )

        current_row = 2

        for ho_ten in sorted(
            thong_ke
        ):
            values = [
                ho_ten,
                thong_ke[ho_ten][0],
                thong_ke[ho_ten][1]
            ]

            for col, value in enumerate(
                values,
                1
            ):
                cell = sheet.cell(
                    current_row,
                    col,
                    value
                )

                cell.font = styles[
                    "font_normal"
                ]

                cell.border = styles[
                    "border"
                ]

                if col == 3:
                    cell.number_format = (
                        styles["money_format"]
                    )

                    cell.alignment = (
                        styles["right"]
                    )
                elif col == 2:
                    cell.alignment = (
                        styles["center"]
                    )
                else:
                    cell.alignment = (
                        styles["left"]
                    )

            BaoCaoController._ap_dung_mau_nguoi(
                sheet,
                current_row,
                1,
                len(headers),
                ho_ten,
                0,
                styles
            )

            current_row += 1

        BaoCaoController._set_excel_widths(
            sheet,
            [
                30,
                18,
                22
            ]
        )

        sheet.freeze_panes = "A2"

        if current_row > 2:
            sheet.auto_filter.ref = (
                f"A1:C{current_row - 1}"
            )

        return sheet

    @staticmethod
    def _tao_sheet_tong_quan(
        workbook,
        du_lieu,
        tu_ngay,
        den_ngay,
        ten_nguoi,
        styles
    ):
        sheet = workbook.create_sheet(
            "Tổng quan"
        )

        BaoCaoController._set_excel_title(
            sheet,
            1,
            "BÁO CÁO QUẢN LÝ CHẤM CƠM",
            2,
            styles
        )

        du_lieu_theo_ngay = {}

        for row in du_lieu:
            du_lieu_theo_ngay.setdefault(
                row[0],
                []
            ).append(row)

        tong_luot_an = sum(
            1
            for row in du_lieu
            if row[3] == 1
        )

        tong_tien = sum(
            (row[4] or 0)
            for row in du_lieu
            if row[3] == 1
        )

        tong_quan = [
            [
                "Từ ngày",
                BaoCaoController._format_ngay(
                    tu_ngay
                )
            ],
            [
                "Đến ngày",
                BaoCaoController._format_ngay(
                    den_ngay
                )
            ],
            [
                "Số ngày",
                len(du_lieu_theo_ngay)
            ],
            [
                "Tổng lượt ăn",
                tong_luot_an
            ],
            [
                "Tổng tiền",
                tong_tien
            ],
            [
                "Trung bình/lượt",
                (
                    tong_tien / tong_luot_an
                    if tong_luot_an
                    else 0
                )
            ]
        ]

        if ten_nguoi.strip():
            tong_quan.insert(
                2,
                [
                    "Người tìm kiếm",
                    ten_nguoi.strip()
                ]
            )

        row_index = 3

        for label, value in tong_quan:
            label_cell = sheet.cell(
                row_index,
                1,
                label
            )

            value_cell = sheet.cell(
                row_index,
                2,
                value
            )

            label_cell.font = styles[
                "font_bold"
            ]

            label_cell.fill = styles[
                "header_fill"
            ]

            label_cell.border = styles[
                "border"
            ]

            value_cell.font = styles[
                "font_normal"
            ]

            value_cell.border = styles[
                "border"
            ]

            if label in [
                "Tổng tiền",
                "Trung bình/lượt"
            ]:
                value_cell.number_format = (
                    styles["money_format"]
                )

                value_cell.alignment = (
                    styles["right"]
                )
            else:
                value_cell.alignment = (
                    styles["left"]
                )

            row_index += 1

        BaoCaoController._set_excel_widths(
            sheet,
            [
                28,
                30
            ]
        )

        return sheet

    @staticmethod
    def _tao_sheet_tuan(
        workbook,
        tu_ngay,
        week_number,
        styles,
        bo_phan_id=None
    ):
        result = (
            BaoCaoController
            .lay_bao_cao_tuan(
                tu_ngay,
                bo_phan_id
            )
        )

        if not result["success"]:
            return None

        sheet = workbook.create_sheet(
            f"Đăng ký tuần {week_number:02d}"
        )

        week_rows = result["data"]

        week_start = result[
            "tu_ngay"
        ]

        week_end = result[
            "den_ngay"
        ]

        headers = [
            "STT",
            "Họ tên",
            "Bộ phận",
            "Thứ 2",
            "Thứ 3",
            "Thứ 4",
            "Thứ 5",
            "Thứ 6",
            "Thứ 7",
            "Chủ nhật",
            "Tổng bữa",
            "Đơn giá",
            "Thành tiền",
            "Đã đóng",
            "Còn nợ"
        ]

        BaoCaoController._set_excel_title(
            sheet,
            1,
            "BÁO CÁO ĐĂNG KÝ ĂN THEO TUẦN",
            len(headers),
            styles
        )

        sheet.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=len(headers)
        )

        info_cell = sheet.cell(
            2,
            1
        )

        info_cell.value = (
            f"Tuần: "
            f"{week_start:%d/%m/%Y}"
            f" - "
            f"{week_end:%d/%m/%Y}"
        )

        info_cell.font = styles[
            "font_normal"
        ]

        info_cell.alignment = styles[
            "center"
        ]

        for col in range(
            1,
            len(headers) + 1
        ):
            sheet.cell(
                2,
                col
            ).border = styles[
                "border"
            ]

        week_days = [
            week_start + timedelta(days=i)
            for i in range(7)
        ]

        date_row = 4
        weekday_row = 5

        for col, header in enumerate(
            headers,
            1
        ):
            cell = sheet.cell(
                weekday_row,
                col
            )

            cell.value = header
            cell.font = styles[
                "font_bold"
            ]

            cell.fill = styles[
                "header_fill"
            ]

            cell.alignment = styles[
                "center"
            ]

            cell.border = styles[
                "border"
            ]

        for index, day in enumerate(
            week_days,
            start=4
        ):
            cell = sheet.cell(
                date_row,
                index
            )

            cell.value = day.strftime(
                "%d/%m/%Y"
            )

            cell.font = styles[
                "font_bold"
            ]

            cell.fill = styles[
                "header_fill"
            ]

            cell.alignment = styles[
                "center"
            ]

            cell.border = styles[
                "border"
            ]

        for col in [
            1,
            2,
            3,
            11,
            12,
            13,
            14,
            15
        ]:
            cell = sheet.cell(
                date_row,
                col
            )

            cell.fill = styles[
                "header_fill"
            ]

            cell.border = styles[
                "border"
            ]

        sheet.row_dimensions[
            date_row
        ].height = 24

        sheet.row_dimensions[
            weekday_row
        ].height = 30

        data_start_row = 6

        for stt, item in enumerate(
            week_rows,
            1
        ):
            values = [
                stt,
                item["HoTen"],
                item["BoPhan"],
                *item["Marks"],
                item["TongBua"],
                item["DonGia"],
                item["ThanhTien"],
                item["DaDong"],
                item["ConNo"]
            ]

            excel_row = (
                data_start_row
                + stt
                - 1
            )

            for col, value in enumerate(
                values,
                1
            ):
                cell = sheet.cell(
                    excel_row,
                    col,
                    value
                )

                cell.font = styles[
                    "font_normal"
                ]

                cell.border = styles[
                    "border"
                ]

                cell.alignment = styles[
                    "center"
                ]

                if col in [
                    2,
                    3
                ]:
                    cell.alignment = styles[
                        "left"
                    ]

                if col in [
                    12,
                    13,
                    14,
                    15
                ]:
                    cell.number_format = (
                        styles["money_format"]
                    )

                    cell.alignment = (
                        styles["right"]
                    )

                if col in [
                    4,
                    5,
                    6,
                    7,
                    8,
                    9,
                    10
                ]:
                    cell.font = styles[
                        "font_bold"
                    ]

                    # Không dùng màu riêng cho từng ô ngày;
                    # màu của cả dòng đại diện cho người đó.
                    if value in ["☑", "☐"]:
                        cell.font = styles[
                            "font_bold"
                        ]

            BaoCaoController._ap_dung_mau_nguoi(
                sheet,
                excel_row,
                1,
                len(headers),
                item["HoTen"],
                item["ConNo"],
                styles
            )

            sheet.row_dimensions[
                excel_row
            ].height = 23

        total_row = (
            data_start_row
            + len(week_rows)
        )

        sheet.merge_cells(
            start_row=total_row,
            start_column=1,
            end_row=total_row,
            end_column=3
        )

        sheet.cell(
            total_row,
            1,
            "TỔNG"
        )

        BaoCaoController._set_excel_total(
            sheet,
            total_row,
            1,
            len(headers),
            styles
        )

        total_meals = sum(
            item["TongBua"]
            for item in week_rows
        )

        total_due = sum(
            item["ThanhTien"]
            for item in week_rows
        )

        total_paid = sum(
            item["DaDong"]
            for item in week_rows
        )

        total_debt = sum(
            item["ConNo"]
            for item in week_rows
        )

        sheet.cell(
            total_row,
            11,
            total_meals
        )

        sheet.cell(
            total_row,
            13,
            total_due
        )

        sheet.cell(
            total_row,
            14,
            total_paid
        )

        sheet.cell(
            total_row,
            15,
            total_debt
        )

        for col in [
            13,
            14,
            15
        ]:
            sheet.cell(
                total_row,
                col
            ).number_format = styles[
                "money_format"
            ]

            sheet.cell(
                total_row,
                col
            ).alignment = styles[
                "right"
            ]

        BaoCaoController._set_excel_widths(
            sheet,
            [
                7,
                28,
                24,
                14,
                14,
                14,
                14,
                14,
                14,
                14,
                14,
                12,
                16,
                18,
                16,
                16
            ]
        )

        sheet.freeze_panes = "D6"

        if total_row > 6:
            sheet.auto_filter.ref = (
                f"A5:O{total_row - 1}"
            )

        sheet.print_title_rows = "1:5"

        return sheet

    @staticmethod
    def _tao_sheet_thang(
        workbook,
        nam,
        thang,
        styles,
        bo_phan_id=None
    ):
        result = (
            BaoCaoController
            .lay_bao_cao_thang(
                nam,
                thang,
                bo_phan_id
            )
        )

        if not result["success"]:
            return None

        sheet = workbook.create_sheet(
            f"Tổng hợp tháng {thang:02d}-{nam}"
        )

        rows = result["data"]

        headers = [
            "STT",
            "Họ tên",
            "Bộ phận",
            "Tổng bữa",
            "Đơn giá",
            "Thành tiền",
            "Đã đóng",
            "Còn nợ",
            "Nộp thừa",
            "Trạng thái"
        ]

        BaoCaoController._set_excel_title(
            sheet,
            1,
            "BÁO CÁO TỔNG HỢP SUẤT ĂN THEO THÁNG",
            len(headers),
            styles
        )

        sheet.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=len(headers)
        )

        info_cell = sheet.cell(
            2,
            1
        )

        info_cell.value = (
            f"Tháng {thang:02d}/{nam}"
        )

        info_cell.font = styles[
            "font_normal"
        ]

        info_cell.alignment = styles[
            "center"
        ]

        for col in range(
            1,
            len(headers) + 1
        ):
            sheet.cell(
                2,
                col
            ).border = styles[
                "border"
            ]

        BaoCaoController._set_excel_header(
            sheet,
            4,
            headers,
            styles
        )

        data_start_row = 5

        for stt, item in enumerate(
            rows,
            1
        ):
            values = [
                stt,
                item["HoTen"],
                item["BoPhan"],
                item["TongBua"],
                item["DonGia"],
                item["ThanhTien"],
                item["DaDong"],
                item["ConNo"],
                item["NopThua"],
                item["TrangThai"]
            ]

            excel_row = (
                data_start_row
                + stt
                - 1
            )

            for col, value in enumerate(
                values,
                1
            ):
                cell = sheet.cell(
                    excel_row,
                    col,
                    value
                )

                cell.font = styles[
                    "font_normal"
                ]

                cell.border = styles[
                    "border"
                ]

                cell.alignment = styles[
                    "center"
                ]

                if col in [
                    2,
                    3
                ]:
                    cell.alignment = styles[
                        "left"
                    ]

                if col in [
                    5,
                    6,
                    7,
                    8,
                    9
                ]:
                    cell.number_format = (
                        styles["money_format"]
                    )

                    cell.alignment = (
                        styles["right"]
                    )

            sheet.row_dimensions[
                excel_row
            ].height = 23

        total_row = (
            data_start_row
            + len(rows)
        )

        sheet.merge_cells(
            start_row=total_row,
            start_column=1,
            end_row=total_row,
            end_column=3
        )

        sheet.cell(
            total_row,
            1,
            "TỔNG"
        )

        BaoCaoController._set_excel_total(
            sheet,
            total_row,
            1,
            len(headers),
            styles
        )

        total_meals = sum(
            item["TongBua"]
            for item in rows
        )

        total_due = sum(
            item["PhaiTra"]
            for item in rows
        )

        total_paid = sum(
            item["DaNop"]
            for item in rows
        )

        total_debt = sum(
            item["ConNo"]
            for item in rows
        )

        total_over = sum(
            item["NopThua"]
            for item in rows
        )

        sheet.cell(
            total_row,
            4,
            total_meals
        )

        sheet.cell(
            total_row,
            6,
            total_due
        )

        sheet.cell(
            total_row,
            7,
            total_paid
        )

        sheet.cell(
            total_row,
            8,
            total_debt
        )

        sheet.cell(
            total_row,
            9,
            total_over
        )

        for col in [
            6,
            7,
            8,
            9
        ]:
            sheet.cell(
                total_row,
                col
            ).number_format = styles[
                "money_format"
            ]

            sheet.cell(
                total_row,
                col
            ).alignment = styles[
                "right"
            ]

        BaoCaoController._set_excel_widths(
            sheet,
            [
                7,
                28,
                24,
                13,
                16,
                18,
                16,
                16,
                16,
                16
            ]
        )

        sheet.freeze_panes = "A5"

        if total_row > 5:
            sheet.auto_filter.ref = (
                f"A4:J{total_row - 1}"
            )

        sheet.print_title_rows = "1:4"

        return sheet

    @staticmethod
    def _tao_sheet_thanh_toan(
        workbook,
        du_lieu,
        styles,
        ten_sheet="Thanh toán"
    ):
        sheet = workbook.create_sheet(
            ten_sheet
        )

        headers = [
            "STT",
            "Họ tên",
            "Bộ phận",
            "Nợ đầu kỳ",
            "Phải trả",
            "Đã nộp",
            "Còn nợ",
            "Nộp thừa",
            "Số dư",
            "Trạng thái"
        ]

        BaoCaoController._set_excel_title(
            sheet,
            1,
            "TÌNH TRẠNG THANH TOÁN",
            len(headers),
            styles
        )

        BaoCaoController._set_excel_header(
            sheet,
            3,
            headers,
            styles
        )

        for stt, item in enumerate(
            du_lieu,
            1
        ):
            row = stt + 3

            values = [
                stt,
                item["HoTen"],
                item["BoPhan"],
                item["NoDauKy"],
                item["PhaiTra"],
                item["DaNop"],
                item["ConNo"],
                item["NopThua"],
                item["SoDuCuoiKy"],
                item["TrangThai"]
            ]

            for col, value in enumerate(
                values,
                1
            ):
                cell = sheet.cell(
                    row,
                    col,
                    value
                )

                cell.font = styles[
                    "font_normal"
                ]

                cell.border = styles[
                    "border"
                ]

                cell.alignment = styles[
                    "center"
                ]

                if col in [
                    2,
                    3
                ]:
                    cell.alignment = styles[
                        "left"
                    ]

                if col in [
                    4,
                    5,
                    6,
                    7,
                    8,
                    9,
                    10
                ]:
                    cell.number_format = (
                        styles["money_format"]
                    )

                    cell.alignment = (
                        styles["right"]
                    )

            BaoCaoController._ap_dung_mau_nguoi(
                sheet,
                row,
                1,
                len(headers),
                item["HoTen"],
                item["ConNo"],
                styles
            )

        total_row = (
            len(du_lieu)
            + 4
        )

        sheet.merge_cells(
            start_row=total_row,
            start_column=1,
            end_row=total_row,
            end_column=3
        )

        sheet.cell(
            total_row,
            1,
            "TỔNG"
        )

        BaoCaoController._set_excel_total(
            sheet,
            total_row,
            1,
            len(headers),
            styles
        )

        totals = {
            4: sum(
                item["NoDauKy"]
                for item in du_lieu
            ),
            5: sum(
                item["PhaiTra"]
                for item in du_lieu
            ),
            6: sum(
                item["DaNop"]
                for item in du_lieu
            ),
            7: sum(
                item["ConNo"]
                for item in du_lieu
            ),
            8: sum(
                item["NopThua"]
                for item in du_lieu
            ),
            9: sum(
                item["SoDuCuoiKy"]
                for item in du_lieu
            )
        }

        for col, value in totals.items():
            cell = sheet.cell(
                total_row,
                col,
                value
            )

            cell.number_format = styles[
                "money_format"
            ]

            cell.alignment = styles[
                "right"
            ]

        BaoCaoController._set_excel_widths(
            sheet,
            [
                7,
                28,
                24,
                18,
                18,
                18,
                18,
                18,
                18,
                16
            ]
        )

        sheet.freeze_panes = "A4"

        if total_row > 4:
            sheet.auto_filter.ref = (
                f"A3:J{total_row - 1}"
            )

        return sheet

    @staticmethod
    def xuat_excel_tuan(
        tu_ngay,
        file_path,
        bo_phan_id=None,
        ten_bo_phan="Tất cả"
    ):
        result = (
            BaoCaoController
            .lay_bao_cao_tuan(
                tu_ngay,
                bo_phan_id
            )
        )

        if not result["success"]:
            return result

        try:
            file_path = (
                BaoCaoController
                ._duong_dan_report(
                    file_path
                )
            )

            wb = Workbook()

            styles = (
                BaoCaoController
                ._tao_style_excel()
            )

            sheet = wb.active
            sheet.title = "Đăng ký tuần"

            rows = result["data"]

            start = result[
                "tu_ngay"
            ]

            end = result[
                "den_ngay"
            ]

            headers = [
                "STT",
                "Họ tên",
                "Bộ phận",
                "Thứ 2",
                "Thứ 3",
                "Thứ 4",
                "Thứ 5",
                "Thứ 6",
                "Thứ 7",
                "Chủ nhật",
                "Tổng bữa",
                "Đơn giá",
                "Thành tiền",
                "Đã đóng",
                "Còn nợ"
            ]

            BaoCaoController._set_excel_title(
                sheet,
                1,
                "BÁO CÁO ĐĂNG KÝ ĂN THEO TUẦN",
                len(headers),
                styles
            )

            sheet.merge_cells(
                start_row=2,
                start_column=1,
                end_row=2,
                end_column=len(headers)
            )

            info_cell = sheet.cell(
                2,
                1
            )

            info_cell.value = (
                f"Tuần: "
                f"{start:%d/%m/%Y}"
                f" - "
                f"{end:%d/%m/%Y}"
                f" | Bộ phận: "
                f"{ten_bo_phan}"
            )

            info_cell.font = styles[
                "font_normal"
            ]

            info_cell.alignment = styles[
                "center"
            ]

            week_days = [
                start + timedelta(days=i)
                for i in range(7)
            ]

            for col, header in enumerate(
                headers,
                1
            ):
                cell = sheet.cell(
                    5,
                    col
                )

                cell.value = header
                cell.font = styles[
                    "font_bold"
                ]

                cell.fill = styles[
                    "header_fill"
                ]

                cell.alignment = styles[
                    "center"
                ]

                cell.border = styles[
                    "border"
                ]

            for index, day in enumerate(
                week_days,
                start=4
            ):
                cell = sheet.cell(
                    4,
                    index
                )

                cell.value = day.strftime(
                    "%d/%m/%Y"
                )

                cell.font = styles[
                    "font_bold"
                ]

                cell.fill = styles[
                    "header_fill"
                ]

                cell.alignment = styles[
                    "center"
                ]

                cell.border = styles[
                    "border"
                ]

            for col in [
                1,
                2,
                3,
                11,
                12,
                13,
                14,
                15
            ]:
                sheet.cell(
                    4,
                    col
                ).fill = styles[
                    "header_fill"
                ]

                sheet.cell(
                    4,
                    col
                ).border = styles[
                    "border"
                ]

            data_start_row = 6

            for stt, item in enumerate(
                rows,
                1
            ):
                values = [
                    stt,
                    item["HoTen"],
                    item["BoPhan"],
                    *item["Marks"],
                    item["TongBua"],
                    item["DonGia"],
                    item["ThanhTien"],
                    item["DaDong"],
                    item["ConNo"]
                ]

                excel_row = (
                    data_start_row
                    + stt
                    - 1
                )

                for col, value in enumerate(
                    values,
                    1
                ):
                    cell = sheet.cell(
                        excel_row,
                        col,
                        value
                    )

                    cell.font = styles[
                        "font_normal"
                    ]

                    cell.border = styles[
                        "border"
                    ]

                    cell.alignment = styles[
                        "center"
                    ]

                    if col in [
                        2,
                        3
                    ]:
                        cell.alignment = (
                            styles["left"]
                        )

                    if col in [
                        12,
                        13,
                        14,
                        15
                    ]:
                        cell.number_format = (
                            styles["money_format"]
                        )

                        cell.alignment = (
                            styles["right"]
                        )

                    if col in [
                        4,
                        5,
                        6,
                        7,
                        8,
                        9,
                        10
                    ]:
                        cell.font = styles[
                            "font_bold"
                        ]

            BaoCaoController._ap_dung_mau_nguoi(
                sheet,
                excel_row,
                1,
                len(headers),
                item["HoTen"],
                item["ConNo"],
                styles
            )

            total_row = (
                data_start_row
                + len(rows)
            )

            sheet.merge_cells(
                start_row=total_row,
                start_column=1,
                end_row=total_row,
                end_column=3
            )

            sheet.cell(
                total_row,
                1,
                "TỔNG"
            )

            BaoCaoController._set_excel_total(
                sheet,
                total_row,
                1,
                len(headers),
                styles
            )

            total_meals = sum(
                x["TongBua"]
                for x in rows
            )

            total_due = sum(
                x["ThanhTien"]
                for x in rows
            )

            total_paid = sum(
                x["DaDong"]
                for x in rows
            )

            total_debt = sum(
                x["ConNo"]
                for x in rows
            )

            sheet.cell(
                total_row,
                11,
                total_meals
            )

            sheet.cell(
                total_row,
                13,
                total_due
            )

            sheet.cell(
                total_row,
                14,
                total_paid
            )

            sheet.cell(
                total_row,
                15,
                total_debt
            )

            for col in [
                13,
                14,
                15
            ]:
                sheet.cell(
                    total_row,
                    col
                ).number_format = styles[
                    "money_format"
                ]

                sheet.cell(
                    total_row,
                    col
                ).alignment = styles[
                    "right"
                ]

            BaoCaoController._set_excel_widths(
                sheet,
                [
                    7,
                    28,
                    24,
                    14,
                    14,
                    14,
                    14,
                    14,
                    14,
                    14,
                    12,
                    16,
                    18,
                    16,
                    16
                ]
            )

            sheet.freeze_panes = "D6"

            if total_row > 6:
                sheet.auto_filter.ref = (
                    f"A5:O{total_row - 1}"
                )

            sheet.print_title_rows = "1:5"

            BaoCaoController._finalize_excel_sheet(
                sheet
            )

            payment_sheet = (
                BaoCaoController
                ._tao_sheet_thanh_toan(
                    wb,
                    rows,
                    styles,
                    "Thanh toán"
                )
            )

            BaoCaoController._finalize_excel_sheet(
                payment_sheet
            )

            wb.save(file_path)

            return {
                "success": True,
                "message": (
                    "Xuất báo cáo tuần thành công."
                ),
                "file_path": str(file_path)
            }

        except Exception as error:
            return {
                "success": False,
                "message": (
                    f"Không thể xuất báo cáo tuần: {error}"
                )
            }

    @staticmethod
    def xuat_excel_thang(
        nam,
        thang,
        file_path,
        bo_phan_id=None,
        ten_bo_phan="Tất cả"
    ):
        result = (
            BaoCaoController
            .lay_bao_cao_thang(
                nam,
                thang,
                bo_phan_id
            )
        )

        if not result["success"]:
            return result

        try:
            file_path = (
                BaoCaoController
                ._duong_dan_report(
                    file_path
                )
            )

            wb = Workbook()

            styles = (
                BaoCaoController
                ._tao_style_excel()
            )

            sheet = wb.active
            sheet.title = "Tổng hợp tháng"

            rows = result["data"]

            headers = [
                "STT",
                "Họ tên",
                "Bộ phận",
                "Tổng bữa",
                "Đơn giá",
                "Thành tiền",
                "Đã đóng",
                "Còn nợ",
                "Nộp thừa",
                "Trạng thái"
            ]

            BaoCaoController._set_excel_title(
                sheet,
                1,
                "BÁO CÁO TỔNG HỢP SUẤT ĂN THEO THÁNG",
                len(headers),
                styles
            )

            sheet.merge_cells(
                start_row=2,
                start_column=1,
                end_row=2,
                end_column=len(headers)
            )

            info_cell = sheet.cell(
                2,
                1
            )

            info_cell.value = (
                f"Tháng "
                f"{int(thang):02d}/{int(nam)}"
                f" | Bộ phận: "
                f"{ten_bo_phan}"
            )

            info_cell.font = styles[
                "font_normal"
            ]

            info_cell.alignment = styles[
                "center"
            ]

            BaoCaoController._set_excel_header(
                sheet,
                4,
                headers,
                styles
            )

            data_start_row = 5

            for stt, item in enumerate(
                rows,
                1
            ):
                values = [
                    stt,
                    item["HoTen"],
                    item["BoPhan"],
                    item["TongBua"],
                    item["DonGia"],
                    item["ThanhTien"],
                    item["DaDong"],
                    item["ConNo"],
                    item["NopThua"],
                    item["TrangThai"]
                ]

                excel_row = (
                    data_start_row
                    + stt
                    - 1
                )

                for col, value in enumerate(
                    values,
                    1
                ):
                    cell = sheet.cell(
                        excel_row,
                        col,
                        value
                    )

                    cell.font = styles[
                        "font_normal"
                    ]

                    cell.border = styles[
                        "border"
                    ]

                    cell.alignment = styles[
                        "center"
                    ]

                    if col in [
                        2,
                        3
                    ]:
                        cell.alignment = (
                            styles["left"]
                        )

                    if col in [
                        5,
                        6,
                        7,
                        8,
                        9
                    ]:
                        cell.number_format = (
                            styles["money_format"]
                        )

                        cell.alignment = (
                            styles["right"]
                        )

            BaoCaoController._ap_dung_mau_nguoi(
                sheet,
                excel_row,
                1,
                len(headers),
                item["HoTen"],
                item["ConNo"],
                styles
            )

            total_row = (
                data_start_row
                + len(rows)
            )

            sheet.merge_cells(
                start_row=total_row,
                start_column=1,
                end_row=total_row,
                end_column=3
            )

            sheet.cell(
                total_row,
                1,
                "TỔNG"
            )

            BaoCaoController._set_excel_total(
                sheet,
                total_row,
                1,
                len(headers),
                styles
            )

            total_meals = sum(
                x["TongBua"]
                for x in rows
            )

            total_due = sum(
                x["PhaiTra"]
                for x in rows
            )

            total_paid = sum(
                x["DaNop"]
                for x in rows
            )

            total_debt = sum(
                x["ConNo"]
                for x in rows
            )

            total_over = sum(
                x["NopThua"]
                for x in rows
            )

            sheet.cell(
                total_row,
                4,
                total_meals
            )

            sheet.cell(
                total_row,
                6,
                total_due
            )

            sheet.cell(
                total_row,
                7,
                total_paid
            )

            sheet.cell(
                total_row,
                8,
                total_debt
            )

            sheet.cell(
                total_row,
                9,
                total_over
            )

            for col in [
                6,
                7,
                8,
                9
            ]:
                sheet.cell(
                    total_row,
                    col
                ).number_format = styles[
                    "money_format"
                ]

                sheet.cell(
                    total_row,
                    col
                ).alignment = styles[
                    "right"
                ]

            BaoCaoController._set_excel_widths(
                sheet,
                [
                    7,
                    28,
                    24,
                    13,
                    16,
                    18,
                    16,
                    16,
                    16,
                    16
                ]
            )

            sheet.freeze_panes = "A5"

            if total_row > 5:
                sheet.auto_filter.ref = (
                    f"A4:J{total_row - 1}"
                )

            sheet.print_title_rows = "1:4"

            BaoCaoController._finalize_excel_sheet(
                sheet
            )

            payment_sheet = (
                BaoCaoController
                ._tao_sheet_thanh_toan(
                    wb,
                    rows,
                    styles,
                    "Thanh toán"
                )
            )

            BaoCaoController._finalize_excel_sheet(
                payment_sheet
            )

            wb.save(file_path)

            return {
                "success": True,
                "message": (
                    "Xuất báo cáo tháng thành công."
                ),
                "file_path": str(file_path)
            }

        except Exception as error:
            return {
                "success": False,
                "message": (
                    f"Không thể xuất báo cáo tháng: {error}"
                )
            }

    @staticmethod
    def xuat_excel(
        tu_ngay,
        den_ngay,
        file_path,
        ten_nguoi=""
    ):
        result = (
            BaoCaoController
            .lay_du_lieu_bao_cao(
                tu_ngay,
                den_ngay,
                ten_nguoi
            )
        )

        if not result["success"]:
            return result

        try:
            file_path = (
                BaoCaoController
                ._duong_dan_report(
                    file_path
                )
            )

            workbook = Workbook()

            styles = (
                BaoCaoController
                ._tao_style_excel()
            )

            du_lieu = result["data"]

            BaoCaoController._tao_sheet_bao_cao(
                workbook,
                du_lieu,
                styles
            )

            BaoCaoController._tao_sheet_theo_ngay(
                workbook,
                du_lieu,
                styles
            )

            BaoCaoController._tao_sheet_theo_nguoi(
                workbook,
                du_lieu,
                styles
            )

            BaoCaoController._tao_sheet_tong_quan(
                workbook,
                du_lieu,
                tu_ngay,
                den_ngay,
                ten_nguoi,
                styles
            )

            start_date = (
                BaoCaoController
                ._to_date(
                    tu_ngay
                )
            )

            end_date = (
                BaoCaoController
                ._to_date(
                    den_ngay
                )
            )

            weeks = (
                BaoCaoController
                ._lay_cac_tuan(
                    start_date,
                    end_date
                )
            )

            for index, week in enumerate(
                weeks,
                1
            ):
                BaoCaoController._tao_sheet_tuan(
                    workbook,
                    week["tu_ngay"],
                    index,
                    styles
                )

            months = (
                BaoCaoController
                ._lay_cac_thang(
                    start_date,
                    end_date
                )
            )

            for item in months:
                BaoCaoController._tao_sheet_thang(
                    workbook,
                    item["nam"],
                    item["thang"],
                    styles
                )

            payment_result = (
                BaoCaoController
                .lay_tong_quan_thanh_toan(
                    start_date,
                    end_date
                )
            )

            if payment_result["success"]:
                payment_sheet = (
                    BaoCaoController
                    ._tao_sheet_thanh_toan(
                        workbook,
                        payment_result["data"],
                        styles,
                        "Công nợ"
                    )
                )

                BaoCaoController._finalize_excel_sheet(
                    payment_sheet
                )

            weekly_payment = (
                BaoCaoController
                .lay_bao_cao_tuan_theo_khoang(
                    start_date,
                    end_date
                )
            )

            if weekly_payment["success"]:
                sheet = workbook.create_sheet(
                    "Thanh toán theo tuần"
                )

                headers = [
                    "Tuần",
                    "Từ ngày",
                    "Đến ngày",
                    "Phải trả",
                    "Đã nộp",
                    "Còn nợ kỳ",
                    "Nộp thừa kỳ",
                    "Còn nợ lũy kế",
                    "Nộp thừa lũy kế",
                    "Đã đủ",
                    "Còn nợ người",
                    "Nộp thừa người",
                    "Quyết toán"
                ]

                BaoCaoController._set_excel_header(
                    sheet,
                    1,
                    headers,
                    styles
                )

                for index, item in enumerate(
                    weekly_payment["data"],
                    2
                ):
                    values = [
                        item["Tuan"],
                        item["TuNgay"].strftime(
                            "%d/%m/%Y"
                        ),
                        item["DenNgay"].strftime(
                            "%d/%m/%Y"
                        ),
                        item["PhaiTra"],
                        item["DaNop"],
                        item["ConNoKy"],
                        item["NopThuaKy"],
                        item["ConNo"],
                        item["NopThua"],
                        item["SoDaDu"],
                        item["SoConNo"],
                        item["SoNopThua"],
                        item["TrangThai"]
                    ]

                    for col, value in enumerate(
                        values,
                        1
                    ):
                        cell = sheet.cell(
                            index,
                            col,
                            value
                        )

                        cell.font = styles[
                            "font_normal"
                        ]

                        cell.border = styles[
                            "border"
                        ]

                        cell.alignment = styles[
                            "center"
                        ]

                        if col in [
                            4,
                            5,
                            6,
                            7,
                            8,
                            9
                        ]:
                            cell.number_format = (
                                styles["money_format"]
                            )

                            cell.alignment = (
                                styles["right"]
                            )

                BaoCaoController._set_excel_widths(
                    sheet,
                    [
                        8,
                        14,
                        14,
                        18,
                        18,
                        18,
                        18,
                        18,
                        20,
                        12,
                        18,
                        20,
                        16
                    ]
                )

                sheet.freeze_panes = "A2"

                if weekly_payment["data"]:
                    sheet.auto_filter.ref = (
                        f"A1:M"
                        f"{len(weekly_payment['data']) + 1}"
                    )

                BaoCaoController._finalize_excel_sheet(
                    sheet
                )

            monthly_payment = (
                BaoCaoController
                .lay_bao_cao_thang_theo_khoang(
                    start_date,
                    end_date
                )
            )

            if monthly_payment["success"]:
                sheet = workbook.create_sheet(
                    "Thanh toán theo tháng"
                )

                headers = [
                    "Năm",
                    "Tháng",
                    "Từ ngày",
                    "Đến ngày",
                    "Phải trả",
                    "Đã nộp",
                    "Còn nợ kỳ",
                    "Nộp thừa kỳ",
                    "Còn nợ lũy kế",
                    "Nộp thừa lũy kế",
                    "Đã đủ",
                    "Còn nợ người",
                    "Nộp thừa người",
                    "Quyết toán"
                ]

                BaoCaoController._set_excel_header(
                    sheet,
                    1,
                    headers,
                    styles
                )

                for index, item in enumerate(
                    monthly_payment["data"],
                    2
                ):
                    values = [
                        item["Nam"],
                        item["Thang"],
                        item["TuNgay"].strftime(
                            "%d/%m/%Y"
                        ),
                        item["DenNgay"].strftime(
                            "%d/%m/%Y"
                        ),
                        item["PhaiTra"],
                        item["DaNop"],
                        item["ConNoKy"],
                        item["NopThuaKy"],
                        item["ConNo"],
                        item["NopThua"],
                        item["SoDaDu"],
                        item["SoConNo"],
                        item["SoNopThua"],
                        item["TrangThai"]
                    ]

                    for col, value in enumerate(
                        values,
                        1
                    ):
                        cell = sheet.cell(
                            index,
                            col,
                            value
                        )

                        cell.font = styles[
                            "font_normal"
                        ]

                        cell.border = styles[
                            "border"
                        ]

                        cell.alignment = styles[
                            "center"
                        ]

                        if col in [
                            5,
                            6,
                            7,
                            8,
                            9,
                            10
                        ]:
                            cell.number_format = (
                                styles["money_format"]
                            )

                            cell.alignment = (
                                styles["right"]
                            )

                BaoCaoController._set_excel_widths(
                    sheet,
                    [
                        10,
                        10,
                        14,
                        14,
                        18,
                        18,
                        18,
                        18,
                        18,
                        20,
                        12,
                        18,
                        20,
                        16
                    ]
                )

                sheet.freeze_panes = "A2"

                if monthly_payment["data"]:
                    sheet.auto_filter.ref = (
                        f"A1:N"
                        f"{len(monthly_payment['data']) + 1}"
                    )

                BaoCaoController._finalize_excel_sheet(
                    sheet
                )

            for sheet in workbook.worksheets:
                BaoCaoController._finalize_excel_sheet(
                    sheet
                )

            workbook.save(file_path)

            return {
                "success": True,
                "message": (
                    "Xuất Excel tổng hợp thành công."
                ),
                "file_path": str(file_path)
            }

        except Exception as error:
            return {
                "success": False,
                "message": (
                    f"Không thể xuất Excel: {error}"
                )
            }

    @staticmethod
    def xuat_csv(
        tu_ngay,
        den_ngay,
        file_path,
        ten_nguoi=""
    ):
        result = (
            BaoCaoController
            .lay_du_lieu_bao_cao(
                tu_ngay,
                den_ngay,
                ten_nguoi
            )
        )

        if not result["success"]:
            return result

        try:
            file_path = (
                BaoCaoController
                ._duong_dan_report(
                    file_path
                )
            )

            with open(
                file_path,
                "w",
                newline="",
                encoding="utf-8-sig"
            ) as file:

                writer = csv.writer(
                    file
                )

                writer.writerow([
                    "Ngày",
                    "Họ tên",
                    "SĐT",
                    "Đã ăn",
                    "Số tiền phải trả",
                    "Ghi chú"
                ])

                for row in result["data"]:
                    writer.writerow([
                        BaoCaoController._format_ngay(
                            row[0]
                        ),
                        row[1],
                        row[2] or "",
                        "Có"
                        if row[3] == 1
                        else "Không",
                        row[4] or 0,
                        row[5] or ""
                    ])

            return {
                "success": True,
                "message": "Xuất CSV thành công.",
                "file_path": str(file_path)
            }

        except Exception as error:
            return {
                "success": False,
                "message": (
                    f"Không thể xuất CSV: {error}"
                )
            }

    @staticmethod
    def tao_dataframe_tong_hop(
        tu_ngay,
        den_ngay,
        bo_phan_id=None
    ):
        result = (
            BaoCaoController
            .lay_tong_quan_thanh_toan(
                tu_ngay,
                den_ngay,
                bo_phan_id
            )
        )

        if not result["success"]:
            return result

        rows = []

        for item in result["data"]:
            rows.append({
                "Người ID": item["NguoiId"],
                "Người": item["HoTen"],
                "Bộ phận": item["BoPhan"],
                "Nợ đầu kỳ": item["NoDauKy"],
                "Phải trả": item["PhaiTra"],
                "Đã nộp": item["DaNop"],
                "Còn nợ": item["ConNo"],
                "Nộp thừa": item["NopThua"],
                "Số dư": item["SoDuCuoiKy"],
                "Trạng thái": item["TrangThai"]
            })

        return {
            "success": True,
            "data": rows,
            "tong_quan": result["tong_quan"]
        }

    @staticmethod
    def tao_dataframe_tong_hop_tuan(
        tu_ngay,
        den_ngay,
        bo_phan_id=None
    ):
        result = (
            BaoCaoController
            .lay_bao_cao_tuan_theo_khoang(
                tu_ngay,
                den_ngay,
                bo_phan_id
            )
        )

        if not result["success"]:
            return result

        rows = []

        for item in result["data"]:
            rows.append({
                "Tuần": item["Tuan"],
                "Từ ngày": item[
                    "TuNgay"
                ].strftime("%d/%m/%Y"),
                "Đến ngày": item[
                    "DenNgay"
                ].strftime("%d/%m/%Y"),
                "Phải trả": item["PhaiTra"],
                "Đã nộp": item["DaNop"],
                "Còn nợ kỳ": item["ConNoKy"],
                "Nộp thừa kỳ": item["NopThuaKy"],
                "Còn nợ lũy kế": item["ConNo"],
                "Nộp thừa lũy kế": item["NopThua"],
                "Đã đủ": item["SoDaDu"],
                "Còn nợ người": item[
                    "SoConNo"
                ],
                "Nộp thừa người": item[
                    "SoNopThua"
                ],
                "Quyết toán": item[
                    "TrangThai"
                ]
            })

        return {
            "success": True,
            "data": rows
        }

    @staticmethod
    def tao_dataframe_chi_tiet_ca_nhan(
        nguoi_an_id,
        tu_ngay,
        den_ngay
    ):
        person = (
            BaoCaoController
            .lay_nguoi_an_theo_id(
                nguoi_an_id
            )
        )

        if not person:
            return {
                "success": False,
                "message": (
                    "Không tìm thấy người ăn."
                ),
                "data": []
            }

        row = (
            BaoCaoController
            ._tao_dong_thanh_toan(
                person[0],
                person[1],
                person[4],
                tu_ngay,
                den_ngay
            )
        )

        return {
            "success": True,
            "data": [
                {
                    "Người ID": row[
                        "NguoiId"
                    ],
                    "Người": row[
                        "HoTen"
                    ],
                    "Bộ phận": row[
                        "BoPhan"
                    ],
                    "Nợ đầu kỳ": row[
                        "NoDauKy"
                    ],
                    "Phải trả": row[
                        "PhaiTra"
                    ],
                    "Đã nộp": row[
                        "DaNop"
                    ],
                    "Còn nợ": row[
                        "ConNo"
                    ],
                    "Nộp thừa": row[
                        "NopThua"
                    ],
                    "Số dư": row[
                        "SoDuCuoiKy"
                    ],
                    "Trạng thái": row[
                        "TrangThai"
                    ]
                }
            ]
        }

    @staticmethod
    def tao_dataframe_chi_tiet_tien_com(
        tu_ngay,
        den_ngay,
        ten_nguoi=""
    ):
        result = (
            BaoCaoController
            .lay_du_lieu_bao_cao(
                tu_ngay,
                den_ngay,
                ten_nguoi
            )
        )

        if not result["success"]:
            return result

        rows = []

        for row in result["data"]:
            rows.append({
                "Ngày": BaoCaoController._format_ngay(
                    row[0]
                ),
                "Họ tên": row[1],
                "SĐT": row[2] or "",
                "Đã ăn": (
                    "Có"
                    if row[3] == 1
                    else "Không"
                ),
                "Số tiền phải trả": (
                    row[4] or 0
                ),
                "Ghi chú": row[5] or ""
            })

        return {
            "success": True,
            "data": rows
        }

    @staticmethod
    def tao_dataframe_lich_su_nop_tien(
        nguoi_an_id,
        tu_ngay=None,
        den_ngay=None
    ):
        connection = get_connection()

        try:
            cursor = connection.cursor()

            sql = """
                SELECT Id,
                       NguoiAnId,
                       NgayNop,
                       SoTien,
                       HinhThuc,
                       GhiChu,
                       NgayTao
                FROM GiaoDichNopTien
                WHERE NguoiAnId = ?
            """

            params = [
                nguoi_an_id
            ]

            if tu_ngay:
                sql += """
                    AND NgayNop >= ?
                """

                params.append(
                    tu_ngay
                )

            if den_ngay:
                sql += """
                    AND NgayNop <= ?
                """

                params.append(
                    den_ngay
                )

            sql += """
                ORDER BY NgayNop DESC,
                         Id DESC
            """

            cursor.execute(
                sql,
                tuple(params)
            )

            rows = []

            for row in cursor.fetchall():
                rows.append({
                    "ID": row[0],
                    "Người ID": row[1],
                    "Ngày nộp": row[2],
                    "Số tiền": float(
                        row[3] or 0
                    ),
                    "Hình thức": row[4] or "",
                    "Ghi chú": row[5] or "",
                    "Ngày tạo": row[6] or ""
                })

            return {
                "success": True,
                "data": rows
            }

        except Exception as error:
            return {
                "success": False,
                "message": (
                    "Không thể lấy lịch sử "
                    f"nộp tiền: {error}"
                ),
                "data": []
            }

        finally:
            connection.close()

    @staticmethod
    def tao_dataframe_suat_an_phat_sinh(
        tu_ngay,
        den_ngay
    ):
        connection = get_connection()

        try:
            cursor = connection.cursor()

            cursor.execute("""
                SELECT s.Id,
                       a.Ngay,
                       s.HoTen,
                       s.DonVi,
                       s.SoLuong,
                       s.DonGia,
                       s.GhiChu,
                       COALESCE(
                           SUM(g.SoTien),
                           0
                       ) AS DaNop
                FROM SuatAnPhatSinh s
                INNER JOIN NgayAn a
                    ON s.NgayAnId = a.Id
                LEFT JOIN GiaoDichSuatAnPhatSinh g
                    ON g.SuatAnPhatSinhId = s.Id
                WHERE a.Ngay BETWEEN ? AND ?
                GROUP BY
                    s.Id,
                    a.Ngay,
                    s.HoTen,
                    s.DonVi,
                    s.SoLuong,
                    s.DonGia,
                    s.GhiChu
                ORDER BY a.Ngay,
                         s.HoTen
            """, (
                tu_ngay,
                den_ngay
            ))

            rows = []

            for row in cursor.fetchall():
                so_luong = int(
                    row[4] or 0
                )

                don_gia = float(
                    row[5] or 0
                )

                phai_tra = (
                    so_luong * don_gia
                )

                da_nop = float(
                    row[7] or 0
                )

                so_du = (
                    da_nop - phai_tra
                )

                rows.append({
                    "ID": row[0],
                    "Ngày": BaoCaoController._format_ngay(
                        row[1]
                    ),
                    "Họ tên": row[2],
                    "Đơn vị": row[3] or "",
                    "Số lượng": so_luong,
                    "Đơn giá": don_gia,
                    "Phải trả": phai_tra,
                    "Đã nộp": da_nop,
                    "Còn nợ": (
                        abs(so_du)
                        if so_du < -0.01
                        else 0
                    ),
                    "Nộp thừa": (
                        so_du
                        if so_du > 0.01
                        else 0
                    ),
                    "Ghi chú": row[6] or ""
                })

            return {
                "success": True,
                "data": rows
            }

        except Exception as error:
            return {
                "success": False,
                "message": (
                    "Không thể lấy suất ăn "
                    f"phát sinh: {error}"
                ),
                "data": []
            }

        finally:
            connection.close()

    @staticmethod
    def tao_dataframe_van_de(
        tu_ngay,
        den_ngay,
        bo_phan_id=None
    ):
        result = (
            BaoCaoController
            .lay_tong_quan_thanh_toan(
                tu_ngay,
                den_ngay,
                bo_phan_id
            )
        )

        if not result["success"]:
            return result

        rows = []

        for item in result["data"]:

            if item["ConNo"] > 0.01:
                rows.append({
                    "Loại": "Còn nợ",
                    "Người ID": item[
                        "NguoiId"
                    ],
                    "Người": item[
                        "HoTen"
                    ],
                    "Bộ phận": item[
                        "BoPhan"
                    ],
                    "Số tiền": item[
                        "ConNo"
                    ],
                    "Chi tiết": (
                        f"Phải trả "
                        f"{item['PhaiTra']:,.0f} đ, "
                        f"đã nộp "
                        f"{item['DaNop']:,.0f} đ"
                    )
                })

            elif item["NopThua"] > 0.01:
                rows.append({
                    "Loại": "Nộp thừa",
                    "Người ID": item[
                        "NguoiId"
                    ],
                    "Người": item[
                        "HoTen"
                    ],
                    "Bộ phận": item[
                        "BoPhan"
                    ],
                    "Số tiền": item[
                        "NopThua"
                    ],
                    "Chi tiết": (
                        f"Phải trả "
                        f"{item['PhaiTra']:,.0f} đ, "
                        f"đã nộp "
                        f"{item['DaNop']:,.0f} đ"
                    )
                })

        rows.sort(
            key=lambda item: (
                item["Loại"],
                -item["Số tiền"]
            )
        )

        return {
            "success": True,
            "data": rows
        }